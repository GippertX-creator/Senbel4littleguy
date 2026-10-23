"""Template-led lab reports with explicit, unfilled observation slots."""
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from fastapi import HTTPException

FIELDS = {
    'title':'实验名称', 'purpose':'实验目的', 'principle':'实验原理',
    'equipment':'仪器与试剂', 'steps':'实验步骤', 'precautions':'注意事项',
    'data':'实验数据', 'results':'实验结果', 'discussion':'分析与讨论', 'conclusion':'实验结论',
}
STANDARD = [
    ('purpose','一、实验目的和要求（必填）'), ('principle','二、实验内容和原理（必填）'),
    ('equipment','三、主要仪器设备'), ('results','四、实验结果与分析（必填）'), ('discussion','五、讨论、心得'),
]
ALIASES = {
    '实验目的':'purpose','目的':'purpose','实验目的和要求':'purpose','实验目的与要求':'purpose',
    '实验原理':'principle','原理':'principle','实验内容和原理':'principle','实验内容与原理':'principle',
    '仪器与试剂':'equipment','实验器材':'equipment','材料与方法':'equipment','仪器和试剂':'equipment',
    '主要仪器设备':'equipment','实验材料与仪器':'equipment',
    '实验步骤':'steps','实验方法':'steps','操作步骤':'steps','方法与步骤':'steps','操作方法与实验步骤':'steps',
    '注意事项':'precautions','实验数据':'data','数据记录':'data','实验数据记录和处理':'data','实验数据记录与处理':'data',
    '实验结果':'results','结果':'results','实验结果与分析':'results','实验结果和分析':'results',
    '分析与讨论':'discussion','讨论':'discussion','讨论、心得':'discussion','讨论与心得':'discussion','讨论和心得':'discussion',
    '结果分析':'discussion','实验结论':'conclusion','结论':'conclusion',
}
DETAIL_LABELS = {
    '同组学生姓名':'group','学生姓名':'name','姓名':'name','学院':'college','系':'department','专业':'major','学号':'student_id',
    '指导教师':'instructor','指导老师':'instructor','课程名称':'course_name','实验类型':'experiment_type',
    '实验项目名称':'title','实验名称':'title','实验地点':'location','实验日期':'date',
    '成绩':None,'批阅日期':None,'教师（签名）':None,
}


def paragraphs(doc):
    yield from doc.paragraphs
    def cells(table):
        seen=set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen: continue
                seen.add(cell._tc)
                yield from cell.paragraphs
                for child in cell.tables: yield from cells(child)
    for table in doc.tables: yield from cells(table)


def replace_span(paragraph,start,end,value):
    if start==end:
        paragraph.add_run(value)
        return
    offset=0; involved=[]
    for run in paragraph.runs:
        stop=offset+len(run.text)
        if stop>start and offset<end: involved.append((run,offset))
        offset=stop
    if not involved: return
    first,a=involved[0];last,b=involved[-1]
    prefix=first.text[:start-a];suffix=last.text[end-b:]
    first.text=prefix+value+suffix
    for run,_ in involved[1:]: run.text=''


def replace_tokens(paragraph,fields):
    tokens={'{{'+alias+'}}':fields.get(key,'') for key,label in FIELDS.items() for alias in (label,key)}
    matches=list(re.finditer('|'.join(re.escape(token) for token in tokens),paragraph.text))
    for match in reversed(matches):replace_span(paragraph,match.start(),match.end(),tokens[match[0]])


def heading_key(text):
    clean=re.sub(r'^[\s一二三四五六七八九十\d、.．:：()（）\-]+','',text)
    clean=re.sub(r'[（(]\s*(?:必填|选填)\s*[)）]','',clean).rstrip('：: \n')
    return ALIASES.get(clean)


def headings(doc):
    # Subheadings such as (一)实验结果 belong to the enclosing major section.
    candidates=[(p,heading_key(p.text)) for p in doc.paragraphs
                if heading_key(p.text) and not re.match(r'^\s*[（(][一二三四五六七八九十]+[）)]',p.text)]
    numbered=[p for p,_ in candidates if re.match(r'^\s*(?:[一二三四五六七八九十]+[、.．]|\d+[、.．])',p.text)]
    if len(numbered)>=2:
        first=next(i for i,(p,_) in enumerate(candidates) if p is numbered[0])
        candidates=[item for i,item in enumerate(candidates) if i<first or any(item[0] is p for p in numbered)]
    return candidates


def template_layout(template=None):
    if not template: return list(STANDARD)
    if template.suffix.lower()!='.docx':
        raise HTTPException(400,'标准版式需要 DOCX 模板，请将学校模板另存为 Word .docx 后上传。')
    doc=Document(template)
    found=[(key,p.text) for p,key in headings(doc)]
    # Table-only forms are supported without moving their cells into the body.
    if not found:
        found=[(heading_key(p.text),p.text) for p in paragraphs(doc) if heading_key(p.text)]
    if not found:
        text='\n'.join(p.text for p in paragraphs(doc))
        found=[(k,label) for k,label in FIELDS.items() if k!='title' and any('{{'+a+'}}' in text for a in (k,label))]
    if not found:
        raise HTTPException(400,'未识别到模板章节或占位符，无法保证学校格式。本次未生成；请使用下载的标准模板，或在学校模板加入支持的章节标题/占位符。')
    unique=[];seen=set()
    for key,title in found:
        if key not in seen: unique.append((key,title));seen.add(key)
    return unique


def record_rows(record):
    return [[subject]+[f'【填写：{field}】' for field in record['fields']] for subject in record['subjects']]


def figure_text(figure,index):
    return (f'【在此粘贴图 {index}：{figure["title"]}】\n'
            f'图片要求：{figure["instruction"]}\n'
            f'图 {index}  {figure["caption"]}\n'+
            '\n'.join(f'【填写：{field}】' for field in figure['fields'])+
            '\n实际观察：【填写本图真实观察到的结构、颜色、形态与分布；未观察到也如实记录】')


def prepare_report(data,layout):
    keys=[k for k,_ in layout]
    fixed={s['key']:s for s in data['sections']}
    parts={key:[] for key in keys}
    titles=dict(layout)
    def put(key,block,fallback=None):
        target=key if key in parts else fallback
        if target not in parts:
            target=key;parts[target]=[];titles[target]=FIELDS[key];keys.append(target)
        if block.get('key')==target and block['kind']=='text':block['title']=''
        parts[target].append(block)
    for key in ('purpose','principle','equipment','steps','precautions'):
        fallback='principle' if key in ('steps','precautions') and 'principle' in parts else None
        section=fixed[key]
        put(key,{'kind':'text',**section},fallback)
    for record in data['records']:
        put('data',{'kind':'table',**record},'results')
    for index,figure in enumerate(data['figures'],1):
        put('results',{'kind':'figure','index':index,**figure})
    for prose in data['analysis']:
        put('results',{'kind':'text',**prose})
    for prose in data['discussion']:
        put('discussion',{'kind':'text',**prose})
    put('conclusion',{'kind':'text',**data['conclusion']},'discussion')
    data['document_sections']=[{'key':k,'title':titles[k],'blocks':parts[k]} for k in keys]
    data['pending']=['身份与实验信息中未提供的项目','逐项原始数据与观察记录','本次实验图片及拍摄信息','分析、结论中标注的实际结果']
    return data['document_sections']


def block_text(block):
    kind=block['kind']
    if kind=='figure': return figure_text(block,block['index'])
    if kind=='table':
        return '\n'.join([block['title'],' | '.join([block['subject_label']]+block['fields']),
                          *[' | '.join(row) for row in record_rows(block)],block['note']])
    return block['title']+'\n'+block['content']


def font(run,size=12,bold=False,east='宋体'):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run._element.get_or_add_rPr().rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia',east)


def source_text(sources):
    return '材料依据：'+'；'.join(dict.fromkeys(f'{s.get("title",s["material_id"])} · {s["unit"]}' for s in sources))


def add_block(doc,block):
    if block['kind']=='figure':
        table=doc.add_table(rows=1,cols=1);table.style='Table Grid';table_borders(table)
        table.rows[0].height=Cm(4.5)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cell=table.cell(0,0)
        p=cell.paragraphs[0];font(p.add_run(f'【在此粘贴图 {block["index"]}：{block["title"]}】'),12,True)
        font(cell.add_paragraph('图片要求：'+block['instruction']).runs[0],10.5)
        p=doc.add_paragraph(f'图 {block["index"]}  {block["caption"]}');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:font(r,10.5)
        for field in block['fields']: font(doc.add_paragraph(f'【填写：{field}】').runs[0],10.5)
        font(doc.add_paragraph('实际观察：【填写本图真实观察到的结构、颜色、形态与分布；未观察到也如实记录】').runs[0])
    elif block['kind']=='table':
        font(doc.add_paragraph(block['title']).runs[0],12,True)
        table=doc.add_table(rows=1,cols=1+len(block['fields']));table.style='Table Grid';table_borders(table)
        for c,t in zip(table.rows[0].cells,[block['subject_label']]+block['fields']):c.text=t
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for values in record_rows(block):
            for c,t in zip(table.add_row().cells,values):c.text=t
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:font(r,10.5)
        if block['note']:font(doc.add_paragraph(block['note']).runs[0],10.5)
    else:
        if block['title']:font(doc.add_paragraph(block['title']).runs[0],12,True)
        for text in block['content'].splitlines():
            p=doc.add_paragraph(text);p.paragraph_format.line_spacing=1.5
            for run in p.runs:font(run)
    if block.get('sources'):
        font(doc.add_paragraph(source_text(block['sources'])).runs[0],9)


def table_borders(table):
    from docx.oxml.ns import qn
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        element=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','6'),('color','999999')]:element.set(qn('w:'+k),v)
        borders.append(element)
    table._tbl.tblPr.append(borders)


def add_sections(doc,sections):
    for section in sections:
        p=doc.add_paragraph(section['title']);p.paragraph_format.keep_with_next=True
        font(p.add_run(''),14,False,'仿宋_GB2312')
        for r in p.runs:font(r,14,False,'仿宋_GB2312')
        for block in section['blocks']: add_block(doc,block)


def insert_blocks(doc,anchor,blocks):
    scratch=Document()
    for attribute in ('page_width','page_height','left_margin','right_margin','top_margin','bottom_margin'):
        setattr(scratch.sections[0],attribute,getattr(doc.sections[0],attribute))
    for block in blocks:add_block(scratch,block)
    # Use the target document's standard table style name, not a foreign package's relationships.
    for child in list(scratch._element.body):
        if child.tag.endswith('}sectPr'):continue
        clone=deepcopy(child);anchor.addnext(clone);anchor=clone


def clean_cover(doc,details,title,stop=None):
    values={**details,'title':title}
    labels='|'.join(re.escape(k) for k in sorted(DETAIL_LABELS,key=len,reverse=True))
    pattern=re.compile(r'('+labels+r')\s*[:：]')
    cover=[]
    for child in doc._element.body:
        if child is stop:break
        cover.extend(child.xpath('.//w:p') if child.tag.endswith('}tbl') else [child] if child.tag.endswith('}p') else [])
    for element in cover:
        p=Paragraph(element,doc)
        if element.getparent().tag.endswith('}tc') and p.text.strip().rstrip('：: ') in DETAIL_LABELS:
            continue  # Paired label/value cells are filled below; keep the label intact.
        matches=list(pattern.finditer(p.text))
        for index in range(len(matches)-1,-1,-1):
            m=matches[index];end=matches[index+1].start() if index+1<len(matches) else len(p.text)
            key=DETAIL_LABELS[m[1]];value=values.get(key,'') if key else ''
            replace_span(p,m.end(),end,(' '+(value or '______')+'  ') if key else '')
        if not matches and re.fullmatch(r'\s*(?:\d{4})?\s*年\s*(?:\d{1,2})?\s*月\s*(?:\d{1,2})?\s*日\s*',p.text):
            replace_span(p,0,len(p.text),details.get('date') or '年    月    日')
    # Cover label/value tables use separate cells.
    for table in doc.tables:
        if stop is not None and table._element is not stop and table._element not in list(doc._element.body)[:list(doc._element.body).index(stop)]:continue
        for row in table.rows:
            if len(row.cells)<2 or row.cells[0]._tc is row.cells[1]._tc:continue
            label=row.cells[0].text.strip().rstrip('：: ')
            if label in DETAIL_LABELS:
                key=DETAIL_LABELS[label];row.cells[1].text=values.get(key,'') or '____________'


def standard_document(details,title):
    doc=Document();s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(1);s.bottom_margin=Cm(2.54);s.left_margin=s.right_margin=Cm(3.175)
    p=doc.add_paragraph('本科实验报告');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:font(r,16,True)
    doc.add_paragraph()
    for label,key in [('姓名','name'),('学院','college'),('系','department'),('专业','major'),('学号','student_id'),('指导教师','instructor')]:
        p=doc.add_paragraph(f'{label}：{details.get(key) or "____________"}')
        for r in p.runs:font(r,14,False,'仿宋_GB2312')
    doc.add_paragraph(details.get('date') or '年    月    日')
    doc.add_page_break()
    p=doc.add_paragraph('实验报告');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:font(r,26,True,'楷体_GB2312')
    for label,key in [('课程名称','course_name'),('实验类型','experiment_type'),('实验项目名称','title'),('学生姓名','name'),('专业','major'),('学号','student_id'),('同组学生姓名','group'),('指导老师','instructor'),('实验地点','location'),('实验日期','date')]:
        p=doc.add_paragraph(f'{label}：{title if key=="title" else details.get(key) or "____________"}')
        for r in p.runs:font(r,12)
    return doc


def scrub_package(doc):
    # Discard old sample photographs instead of leaving their blobs in the exported ZIP.
    used=set(doc._element.xpath('.//@r:embed|.//@r:link'))
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
    for attr in ('author','last_modified_by','comments','keywords','subject'):
        setattr(doc.core_properties,attr,'')


def export_report(data,output:Path,template:Path|None=None):
    layout=template_layout(template);sections=prepare_report(data,layout)
    details=data.get('details',{})
    warnings=[]
    if template:
        doc=Document(template);anchors=headings(doc)
        if anchors:
            clean_cover(doc,details,data['title'],anchors[0][0]._p)
            fields={s['key']:'\n\n'.join(block_text(b) for b in s['blocks']) for s in sections}
            fields['title']=data['title']
            for child in doc._element.body:
                if child is anchors[0][0]._p:break
                for p in paragraphs_in(child,doc):replace_tokens(p,fields)
            children=list(doc._element.body)
            for index,(p,key) in enumerate(anchors):
                for shape in p._p.xpath('.//w:drawing|.//w:pict|.//w:object'):
                    shape.getparent().remove(shape)
                start=children.index(p._p)+1
                end=children.index(anchors[index+1][0]._p) if index+1<len(anchors) else len(children)
                for child in children[start:end]:
                    if not child.tag.endswith('}sectPr') and child.getparent() is not None:child.getparent().remove(child)
                section=next(s for s in sections if s['key']==key)
                insert_blocks(doc,p._p,section['blocks'])
            present={k for _,k in anchors}
            add_sections(doc,[s for s in sections if s['key'] not in present])
        else:
            stop=next((c for c in doc._element.body if
                       any(heading_key(p.text) or any('{{'+a+'}}' in p.text for k,label in FIELDS.items() if k!='title' for a in (k,label))
                           for p in paragraphs_in(c,doc))),None)
            clean_cover(doc,details,data['title'],stop)
            # Split-run placeholders retain their run style; complex content becomes readable text.
            fields={s['key']:'\n\n'.join(block_text(b) for b in s['blocks']) for s in sections}
            fields.update({k:fields.get(k,'') for k in FIELDS});fields['title']=data['title']
            used=set()
            for p in list(paragraphs(doc)):
                for k,label in FIELDS.items():
                    if any('{{'+a+'}}' in p.text for a in (k,label)):used.add(k)
                replace_tokens(p,fields)
            if used:
                add_sections(doc,[s for s in sections if s['key'] not in used])
                warnings.append('占位符保留原文字版式；图片与数据位置以文字标记呈现。')
            else:
                # Table heading/value layouts: preserve the outer form and populate its value cell.
                used=set()
                for table in doc.tables:
                    for row in table.rows:
                        seen=set()
                        for index,cell in enumerate(row.cells):
                            if cell._tc in seen:continue
                            seen.add(cell._tc)
                            key=heading_key(cell.text)
                            if not key or key in used:continue
                            target=row.cells[index+1] if index+1<len(row.cells) and row.cells[index+1]._tc is not cell._tc else cell
                            if target is not cell:
                                for node in list(target._tc):
                                    if not node.tag.endswith('}tcPr'):target._tc.remove(node)
                                anchor=target.add_paragraph()._p
                            else:anchor=cell.paragraphs[0]._p
                            insert_blocks(doc,anchor,next(s for s in sections if s['key']==key)['blocks']);used.add(key)
                add_sections(doc,[s for s in sections if s['key'] not in used])
        warnings.append('已按学校模板章节填入，清除正文示例与示例图片；请核对分页、图框空间及学校要求。')
    else:
        doc=standard_document(details,data['title']);add_sections(doc,sections)
    warnings.append('仅填写已标注的数据、照片及未知身份信息；分析和结论中的实际结果填空须与原始记录一致。PDF 为阅读版，学校版式请使用 Word，或在 Word/WPS 中另存 PDF。')
    scrub_package(doc);doc.core_properties.title=data['title'];doc.save(output)
    data['format_version']=2
    data['template_sections']=[title for _,title in layout]
    return warnings


def paragraphs_in(element,doc):
    if element.tag.endswith('}p'):yield Paragraph(element,doc)
    elif element.tag.endswith('}tbl'):
        for child in element.xpath('.//w:p'):yield Paragraph(child,doc)
