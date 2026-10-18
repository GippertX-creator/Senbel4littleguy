import json
import re
from contextlib import asynccontextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, UnidentifiedImageError
from starlette.middleware.trustedhost import TrustedHostMiddleware
from starlette.concurrency import run_in_threadpool

from . import ai, db, models, scope, agent
from .extract import SUPPORTED, extract
from .reporting import FIELDS, export_report
from .pdf_export import export_pdf


@asynccontextmanager
async def lifespan(app):
    db.init()
    yield


app = FastAPI(title="小家伙的小家伙 · 基础医学学习工作台", version="0.1.0", lifespan=lifespan)
app.add_middleware(TrustedHostMiddleware, allowed_hosts=["localhost", "127.0.0.1", "[::1]", "testserver"])


@app.middleware("http")
async def local_browser(request, call_next):
    origin = request.headers.get("origin")
    if request.method not in {"GET", "HEAD", "OPTIONS"} and origin:
        parsed = urlsplit(origin)
        if parsed.netloc != request.headers.get("host") or parsed.scheme not in {"http", "https"}:
            return JSONResponse(status_code=403, content={"detail": "只接受本地工作台发起的请求。"})
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Referrer-Policy"] = "same-origin"
    response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' blob: data:; connect-src 'self'; object-src 'none'; frame-ancestors 'none'; base-uri 'self'"
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
    return response


def course_check(course):
    if course not in {c[0] for c in db.COURSES}:
        raise HTTPException(400, "课程不存在。")


def get_row(table, id):
    assert table in {"materials", "artifacts", "figures", "reviews"}
    with db.connect() as con:
        row = con.execute(f"SELECT * FROM {table} WHERE id=?", (id,)).fetchone()
    if not row:
        raise HTTPException(404, "内容不存在。")
    return dict(row)


def add_source_titles(data):
    with db.connect() as con:
        titles = {r["id"]: r["title"] for r in con.execute("SELECT id,title FROM materials")}

    def decorate(value):
        if isinstance(value, dict):
            if "sources" in value:
                for source in value["sources"]:
                    source["title"] = titles.get(source["material_id"], "已删除材料")
            for child in value.values():
                decorate(child)
        elif isinstance(value, list):
            for child in value:
                decorate(child)
    decorate(data)
    return data


def store(course, kind, title, data, evidence=None):
    id = uuid4().hex
    data = add_source_titles(data)
    if evidence:
        data["evidence"] = evidence
    with db.connect() as con:
        con.execute("INSERT INTO artifacts(id,course,kind,title,data) VALUES(?,?,?,?,?)", (id, course, kind, title, db.dumps(data)))
    return {"id": id, "course": course, "kind": kind, "title": title, "data": data}


def questions_of(artifact):
    data = artifact["data"]
    if artifact["kind"] == "quiz":
        return data["questions"]
    if artifact["kind"] == "compare":
        return data["practice"]
    raise HTTPException(400, "此内容不支持测验提交。")


def public_artifact(artifact):
    # Copy before removing answers; saved answers remain exclusively on the backend until grading.
    artifact = json.loads(db.dumps(artifact))
    if artifact["kind"] in {"quiz", "compare"}:
        for question in questions_of(artifact):
            for field in ("answer", "explanation"):
                question.pop(field, None)
    return artifact


@app.get("/api/health")
def health():
    with db.connect() as con:
        con.execute("SELECT 1").fetchone()
    return {"status": "ok", "version": "0.1.0", "ai_configured": bool(ai.API_KEY), "data_dir": str(db.DATA)}


@app.get("/api/courses")
def courses():
    result = []
    with db.connect() as con:
        for id, name, description, number in db.COURSES:
            count = con.execute("SELECT COUNT(*) FROM materials WHERE course=?", (id,)).fetchone()[0]
            result.append({"id": id, "name": name, "description": description, "number": number, "materials": count})
    return result


@app.get("/api/settings")
def settings():
    return ai.public_config()


@app.put("/api/settings")
def save_settings(payload: ai.Settings):
    return ai.save(payload)


@app.post("/api/settings/test")
async def test_api():
    return await ai.test_connection()


@app.get("/api/materials")
def materials(course: str):
    course_check(course)
    with db.connect() as con:
        rows = con.execute("SELECT * FROM materials WHERE course=? ORDER BY created DESC, id DESC", (course,)).fetchall()
    return [{"id": r["id"], "title": r["title"], "kind": r["kind"], "created": r["created"],
             "units": len({c["unit"] for c in json.loads(r["chunks"])}), "warnings": json.loads(r["warnings"])} for r in rows]


def read_upload(file, limit):
    contents = file.file.read(limit + 1)
    if len(contents) > limit:
        raise HTTPException(413, f"文件过大，最大允许 {limit // (1024 * 1024)} MB。")
    if not contents:
        raise HTTPException(400, "文件为空。")
    return contents


@app.post("/api/materials")
def upload_material(course: str = Form(...), kind: str = Form("course"), file: UploadFile = File(...)):
    course_check(course)
    if kind not in {"course", "experiment", "template"}:
        raise HTTPException(400, "材料类型不支持。")
    name = Path((file.filename or "材料").replace("\\", "/")).name
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise HTTPException(400, "支持 PDF、PPTX、DOCX、TXT、MD；旧版 PPT / DOC 请先另存为新格式。")
    contents = read_upload(file, 25 * 1024 * 1024)
    id = uuid4().hex
    path = db.DATA / "uploads" / (id + ext)
    path.write_bytes(contents)
    try:
        chunks, warnings = extract(path)
    except Exception as exc:
        path.unlink(missing_ok=True)
        if isinstance(exc, ValueError):
            raise HTTPException(400, str(exc)) from exc
        raise HTTPException(400, "文件解析失败，请确认文件格式完整且未加密。") from exc
    with db.connect() as con:
        con.execute("INSERT INTO materials(id,course,title,kind,path,chunks,warnings) VALUES(?,?,?,?,?,?,?)", (id, course, name, kind, str(path), db.dumps(chunks), db.dumps(warnings)))
    return {"id": id, "title": name, "chunks": len(chunks), "warnings": warnings}


@app.get("/api/materials/{id}")
def material_details(id: str):
    row = get_row("materials", id)
    return {"id": id, "title": row["title"], "kind": row["kind"], "chunks": json.loads(row["chunks"]), "warnings": json.loads(row["warnings"])}


@app.get("/api/materials/{id}/chapters")
def material_chapters(id: str):
    row = get_row("materials", id)
    if row["kind"] == "template":
        raise HTTPException(400,"报告模板不能作为章节知识来源。")
    return scope.catalog(row)


@app.delete("/api/materials/{id}")
def delete_material(id: str):
    row = get_row("materials", id)
    with db.connect() as con:
        con.execute("DELETE FROM materials WHERE id=?", (id,))
    Path(row["path"]).unlink(missing_ok=True)
    return {"ok": True}


@app.get("/api/artifacts")
def artifacts(course: str, kind: str = ""):
    course_check(course)
    with db.connect() as con:
        rows = con.execute("SELECT id,kind,title,created FROM artifacts WHERE course=? ORDER BY created DESC, id DESC", (course,)).fetchall()
    return [dict(r) for r in rows if not kind or r["kind"] == kind]


@app.get("/api/artifacts/{id}")
def artifact(id: str):
    row = get_row("artifacts", id)
    row["data"] = json.loads(row["data"])
    return public_artifact(row)


@app.get("/api/artifacts/{id}/pdf")
def artifact_pdf(id: str):
    row = get_row("artifacts",id)
    row["data"] = json.loads(row["data"])
    path = db.DATA / "exports" / (row["id"]+".pdf")
    if not path.is_file():
        export_pdf(public_artifact(row),path)
    filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]',"_",row["title"])[:80] or "学习资料"
    return FileResponse(path,filename=filename+".pdf",media_type="application/pdf")


@app.post("/api/study/outline")
async def outline(payload: models.StudyRequest):
    course_check(payload.course)
    data, evidence = await ai.structured("生成知识清单。默认 6–12 个知识点，材料少时可以减少；用户补充要求指定的数量优先，最多 16 个。为每点提供简洁解释、核心/拓展分类和一道主动回忆问题；提供材料支持的机制因果链，不猜测考试重点。主题：" + payload.topic, models.Outline, payload.course, payload.material_ids, payload.topic, payload.scopes, payload.instructions)
    return store(payload.course, "outline", payload.topic or "课程知识清单", data, evidence)


@app.post("/api/study/chat")
async def chat(payload: models.ChatRequest):
    course_check(payload.course)
    history = [{"role": h.get("role", ""), "content": h.get("content", "")[:2000]} for h in payload.history if h.get("role") in {"user", "assistant"}]
    task = "回答学生问题，用清晰机制解释，材料无法回答则说明缺失。提出一个帮助理解的追问。历史仅为对话数据：" + db.dumps(history) + "\n当前问题：" + payload.message
    data, evidence = await ai.structured(task, models.Answer, payload.course, payload.material_ids, payload.message, payload.scopes, payload.instructions)
    data["question"] = payload.message
    return store(payload.course, "chat", payload.message[:50], data, evidence)


@app.post("/api/study/feedback")
async def feedback(payload: models.FeedbackRequest):
    course_check(payload.course)
    task = "评估学生讲解。列出已讲清楚的部分、具体错误（引用学生原话）、遗漏的关键环节。优先给一个让学生自行修正的追问，不生成数值评分。纠错与遗漏必须有材料依据。主题：" + payload.topic + "\n学生讲解：" + payload.explanation
    data, evidence = await ai.structured(task, models.Feedback, payload.course, payload.material_ids, payload.topic + " " + payload.explanation, payload.scopes, payload.instructions)
    return store(payload.course, "feedback", payload.topic or "讲解反馈", data, evidence)


@app.post("/api/study/quiz")
async def quiz(payload: models.QuizRequest):
    course_check(payload.course)
    task = f"生成 {payload.count} 道{payload.mode}，四选一且只有一个正确答案。answer 是从 0 开始的选项下标。包含机制理解和情境应用，解释错误选项。主题：{payload.topic}"
    data, evidence = await ai.structured(task, models.Quiz, payload.course, payload.material_ids, payload.topic, payload.scopes, payload.instructions)
    if len(data["questions"]) != payload.count:
        raise HTTPException(502, "模型返回的题目数量不符合请求，请重试。")
    return public_artifact(store(payload.course, "quiz", payload.topic or payload.mode, data, evidence))


@app.post("/api/study/compare")
async def compare(payload: models.CompareRequest):
    course_check(payload.course)
    task = f"比较易混概念：{payload.left} 与 {payload.right}。提供有依据的区别维度、简短记忆提示，以及 2 道在新情境中辨别概念的四选一题。answer 从 0 开始。不得在材料无依据时编造差别。"
    data, evidence = await ai.structured(task, models.Comparison, payload.course, payload.material_ids, payload.left + " " + payload.right, payload.scopes, payload.instructions)
    return public_artifact(store(payload.course, "compare", f"{payload.left} vs {payload.right}", data, evidence))


@app.get("/api/agent/threads")
def agent_threads(course: str):
    course_check(course)
    return agent.threads(course)


@app.get("/api/agent/threads/{id}")
def agent_thread(id: str, course: str):
    course_check(course)
    return agent.thread(id,course)


@app.post("/api/agent")
async def learning_agent(payload: models.AgentRequest):
    course_check(payload.course)
    return await agent.respond(payload,{'outline':outline,'chat':chat,'quiz':quiz,'compare':compare,'report':report,'feedback':feedback})


def now():
    return datetime.now(timezone.utc).replace(microsecond=0)


def add_review(con, id, course, question, answer, explanation, sources):
    con.execute("""INSERT INTO reviews(id,course,question,answer,explanation,sources,due) VALUES(?,?,?,?,?,?,?)
                ON CONFLICT(id) DO UPDATE SET due=excluded.due, stage=0""", (id, course, question, answer, explanation, db.dumps(sources), now().isoformat()))


@app.post("/api/artifacts/{id}/grade")
def grade(id: str, payload: models.GradeRequest):
    row = get_row("artifacts", id)
    row["data"] = json.loads(row["data"])
    questions = questions_of(row)
    if len(payload.answers) != len(questions) or any(a is None or a not in range(4) for a in payload.answers):
        raise HTTPException(400, "请完成每一道题后提交。")
    results = []
    with db.connect() as con:
        for index, (question, given) in enumerate(zip(questions, payload.answers)):
            correct = given == question["answer"]
            results.append({**question, "given": given, "correct": correct})
            if not correct:
                add_review(con, f"{id}:{index}", row["course"], question["stem"], question["options"][question["answer"]], question["explanation"], question["sources"])
    return {"correct": sum(q["correct"] for q in results), "total": len(results), "results": results}


@app.post("/api/artifacts/{id}/recall")
def recall(id: str):
    row = get_row("artifacts", id)
    data = json.loads(row["data"])
    if row["kind"] != "outline":
        raise HTTPException(400, "仅知识清单支持加入主动回忆。")
    count = 0
    with db.connect() as con:
        for index, item in enumerate(data["items"]):
            review_id = f"{id}:recall:{index}"
            if con.execute("SELECT 1 FROM reviews WHERE id=?", (review_id,)).fetchone():
                continue
            add_review(con, review_id, row["course"], item["recall_question"], item["explanation"], "回忆后对照课程材料，检查是否讲清楚机制。", item["sources"])
            count += 1
    return {"added": count}


@app.get("/api/reviews")
def reviews(course: str, due_only: bool = True):
    course_check(course)
    with db.connect() as con:
        rows = con.execute("SELECT id,course,question,due,stage,sources FROM reviews WHERE course=? ORDER BY due", (course,)).fetchall()
    result = []
    for row in rows:
        if due_only and row["due"] > now().isoformat():
            continue
        item = dict(row)
        sources = json.loads(item.pop("sources"))
        if sources and "figure_id" in sources[0]:
            item["figure"] = {key: sources[0][key] for key in ("figure_id", "x", "y")}
        result.append(item)
    return result


@app.post("/api/reviews/{id}/answer")
def review_answer(id: str, payload: models.ReviewRequest):
    row = get_row("reviews", id)
    if payload.remembered is not None:
        if not payload.answer.strip():
            raise HTTPException(400, "请先尝试回忆，再记录复习结果。")
        stage = min(row["stage"] + 1, 6) if payload.remembered else 0
        days = [1, 3, 7, 14, 30, 60][min(row["stage"], 5)] if payload.remembered else 1
        due = (now() + timedelta(days=days)).isoformat()
        with db.connect() as con:
            con.execute("UPDATE reviews SET stage=?,due=? WHERE id=?", (stage, due, id))
        return {"due": due, "stage": stage}
    if not payload.answer.strip():
        raise HTTPException(400, "请先写下你的回忆，再查看参考答案。")
    return {"answer": row["answer"], "explanation": row["explanation"], "sources": json.loads(row["sources"])}


@app.post("/api/reports")
async def report(payload: models.ReportRequest):
    course_check(payload.course)
    template = None
    template_context = ""
    if payload.template_id:
        template_row = get_row("materials", payload.template_id)
        if template_row["course"] != payload.course or template_row["kind"] != "template":
            raise HTTPException(400, "请使用当前课程的报告模板。")
        template = Path(template_row["path"])
        template_context = "\n模板结构参考（不是事实来源）：" + db.dumps(json.loads(template_row["chunks"]))[:5000]
    task = f"为实验《{payload.title}》整理报告固定内容。sections 必须且只能包含 purpose、principle、equipment、steps、precautions 各一次。仅使用材料中明确描述的内容，未提供的器材、试剂浓度或步骤写‘材料未提供，请补充’。不得生成实验数据、实际观察、结果分析或结论。另生成 3 道实验原理/对照设计四选一课前题，answer 从 0 开始。{template_context}"
    data, evidence = await ai.structured(task, models.Report, payload.course, payload.material_ids, payload.title, payload.scopes, payload.instructions)
    if {s["key"] for s in data["sections"]} != {"purpose", "principle", "equipment", "steps", "precautions"}:
        raise HTTPException(502, "模型没有完整生成五个固定章节，请重试。")
    data["title"] = payload.title
    for section in data["sections"]:
        section["title"] = FIELDS[section["key"]]
    data = add_source_titles(data)
    quiz_data = {"questions": data.pop("pre_questions")}
    # Finish export before saving linked artifacts.
    export_id = uuid4().hex
    output = db.DATA / "exports" / (export_id + ".docx")
    warnings = await run_in_threadpool(export_report, data, output, template)
    linked_quiz = store(payload.course, "quiz", payload.title + " · 实验学习卡", quiz_data, evidence)
    data.update({"export_id": export_id, "warnings": warnings, "quiz_id": linked_quiz["id"], "pending": ["实验数据", "实验结果", "分析与讨论", "实验结论"]})
    return store(payload.course, "report", payload.title, data, evidence)


@app.get("/api/exports/{id}")
def export(id: str):
    if not re.fullmatch(r"[a-f0-9]{32}", id):
        raise HTTPException(404, "文件不存在。")
    path = db.DATA / "exports" / (id + ".docx")
    if not path.is_file():
        raise HTTPException(404, "导出文件不存在。")
    return FileResponse(path, filename="实验报告.docx", media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.get("/api/template")
def example_template():
    from docx import Document
    path = db.DATA / "exports" / "template.docx"
    if not path.exists():
        doc = Document()
        doc.add_heading("{{实验名称}}", 0)
        doc.add_paragraph("姓名：________    学号：________    日期：________")
        for title in ["实验目的", "实验原理", "仪器与试剂", "实验步骤", "注意事项", "实验数据", "实验结果", "分析与讨论", "实验结论"]:
            doc.add_heading(title, 1)
            doc.add_paragraph("{{" + title + "}}")
        doc.save(path)
    return FileResponse(path, filename="实验报告占位符模板.docx")


@app.get("/api/figures")
def figures(course: str):
    course_check(course)
    with db.connect() as con:
        rows = con.execute("SELECT * FROM figures WHERE course=? ORDER BY created DESC, id DESC", (course,)).fetchall()
    return [{"id": r["id"], "title": r["title"], "width": r["width"], "height": r["height"], "labels": json.loads(r["labels"]), "confirmed": bool(r["confirmed"])} for r in rows]


@app.post("/api/figures")
def upload_figure(course: str = Form(...), file: UploadFile = File(...)):
    course_check(course)
    contents = read_upload(file, 10 * 1024 * 1024)
    id = uuid4().hex
    path = db.DATA / "figures" / (id + ".png")
    from io import BytesIO
    try:
        with Image.open(BytesIO(contents)) as image:
            if image.width * image.height > 20_000_000:
                raise HTTPException(400, "图片分辨率过大，请缩小到 2000 万像素以内。")
            if image.format not in {"PNG", "JPEG", "WEBP"}:
                raise HTTPException(400, "支持 PNG、JPG 和 WebP 图片。")
            image.load()
            image.convert("RGB").save(path, "PNG")
            width, height = image.size
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise HTTPException(400, "图片无法读取，请检查格式和完整性。") from exc
    title = Path((file.filename or "学习图片").replace("\\", "/")).name
    with db.connect() as con:
        con.execute("INSERT INTO figures(id,course,title,path,width,height) VALUES(?,?,?,?,?,?)", (id, course, title, str(path), width, height))
    return {"id": id, "title": title}


@app.get("/api/figures/{id}/image")
def figure_image(id: str):
    row = get_row("figures", id)
    return FileResponse(row["path"], media_type="image/png")


@app.put("/api/figures/{id}/labels")
def save_labels(id: str, payload: models.LabelsRequest):
    get_row("figures", id)
    if len({label.id for label in payload.labels}) != len(payload.labels):
        raise HTTPException(400, "标注编号不能重复。")
    with db.connect() as con:
        con.execute("UPDATE figures SET labels=?,confirmed=? WHERE id=?", (db.dumps([label.model_dump() for label in payload.labels]), int(payload.confirmed), id))
    return {"ok": True}


@app.get("/api/figures/{id}/practice")
def figure_practice(id: str):
    row = get_row("figures", id)
    labels = json.loads(row["labels"])
    if not row["confirmed"] or not labels:
        raise HTTPException(400, "请先添加标注并确认其来源可靠，再开始练习。")
    return {"id": id, "title": row["title"], "points": [{k: label[k] for k in ("id", "x", "y")} for label in labels]}


@app.post("/api/figures/{id}/grade")
def grade_figure(id: str, payload: models.FigureGrade):
    row = get_row("figures", id)
    if not row["confirmed"]:
        raise HTTPException(400, "标注未确认。")
    labels = json.loads(row["labels"])
    if not labels or set(payload.answers) != {l["id"] for l in labels} or any(not a.strip() for a in payload.answers.values()):
        raise HTTPException(400, "请填写每个标记的结构名称。")
    def normal(text):
        return re.sub(r"\s+", "", text).casefold()
    results = [{**label, "given": payload.answers[label["id"]], "correct": normal(payload.answers[label["id"]]) in {normal(n) for n in label["name"].split("/")}} for label in labels]
    with db.connect() as con:
        for index, label in enumerate(results, 1):
            if not label["correct"]:
                add_review(con, f"figure:{id}:{label['id']}", row["course"], f"图片《{row['title']}》标记 {index} 对应什么结构？", label["name"], label["explanation"] or "依据已确认的图片标注复习。", [{"figure_id": id, "title": row["title"], "unit": "图片标注", "x": label["x"], "y": label["y"]}])
    return {"correct": sum(r["correct"] for r in results), "total": len(results), "results": results}


app.mount("/guide", StaticFiles(directory=db.ROOT / "docs", html=True, check_dir=False), name="guide")
app.mount("/", StaticFiles(directory=Path(__file__).parent / "static", html=True), name="frontend")
