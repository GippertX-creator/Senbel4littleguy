import re
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

SUPPORTED = {".pdf", ".pptx", ".docx", ".txt", ".md"}


def safe_office(path):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 6000 or sum(i.file_size for i in entries) > 120 * 1024 * 1024:
                raise ValueError("Office 文件展开后过大，请拆分后上传。")
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 DOCX / PPTX 文件。") from exc


def split_text(text, unit):
    text = re.sub(r"[ \t]+", " ", text).strip()
    return [{"unit": unit, "text": text[i:i + 1400]} for i in range(0, len(text), 1300) if text[i:i + 1400].strip()]


def extract(path: Path):
    ext = path.suffix.lower()
    chunks, warnings = [], []
    if ext in {".docx", ".pptx"}:
        safe_office(path)
    if ext == ".pdf":
        with fitz.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("请上传已解除密码保护的 PDF。")
            if len(doc) > 600:
                raise ValueError("PDF 超过 600 页，请按章节拆分。")
            empty = []
            for index, page in enumerate(doc):
                text = page.get_text()
                chunks.extend(split_text(text, f"第 {index + 1} 页"))
                if not text.strip():
                    empty.append(str(index + 1))
            if empty:
                warnings.append("以下 PDF 页未提取到文字（图片/扫描页需要另行提供文字）：" + ", ".join(empty[:30]))
    elif ext == ".pptx":
        presentation = Presentation(path)
        if len(presentation.slides) > 600:
            raise ValueError("PPT 超过 600 页，请拆分。")

        def shape_text(shapes):
            parts = []
            for shape in shapes:
                if hasattr(shape, "shapes"):
                    parts.extend(shape_text(shape.shapes))
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
            return parts

        for index, slide in enumerate(presentation.slides):
            text = "\n".join(shape_text(slide.shapes))
            if slide.has_notes_slide:
                text += "\n" + slide.notes_slide.notes_text_frame.text
            chunks.extend(split_text(text, f"第 {index + 1} 张幻灯片"))
        warnings.append("仅提取幻灯片文字、表格与备注；图像中的文字和图示未自动解析。")
    elif ext == ".docx":
        doc = Document(path)
        paragraph_index = table_index = 0
        # Preserve table placement so chapter boundaries cannot absorb other chapters' tables.
        for block in doc.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_index += 1
                chunks.extend(split_text(block.text, f"第 {paragraph_index} 段"))
            elif isinstance(block, Table):
                table_index += 1
                text = "\n".join(" | ".join(c.text for c in row.cells) for row in block.rows)
                chunks.extend(split_text(text, f"表 {table_index}"))
    elif ext in {".txt", ".md"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise ValueError("文本编码无法识别，请保存为 UTF-8 后上传。") from exc
        for index, part in enumerate(text.split("\n\n")):
            chunks.extend(split_text(part, f"段落 {index + 1}"))
    if len(chunks) > 1800:
        raise ValueError("材料文字量过大，请按章节拆分。")
    if not chunks:
        raise ValueError("未提取到文字。扫描 PDF 请先进行 OCR；图片请使用指图学习。")
    return chunks, warnings
