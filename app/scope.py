"""Local chapter discovery and exact source filtering; no model decides boundaries."""
import json
from pathlib import Path
import re
import unicodedata

from docx import Document
from pptx import Presentation
from fastapi import HTTPException
import fitz

from .extract import extract

HEADING = re.compile(r"^(?:第[一二三四五六七八九十百零〇两\d]+[章节篇部][\s、.：:．-]*.{0,70}|chapter\s+\d+.{0,70})$", re.I)


def chunks_of(row):
    # Older DOCX uploads stored all tables at the end; re-read their original order.
    path = Path(row["path"])
    if path.suffix.lower() == ".docx" and path.is_file():
        raw = extract(path)[0]
    else:
        raw = json.loads(row["chunks"])
    chunks = []
    for chunk in raw:
        text = chunk["text"]
        if chunks and chunks[-1]["unit"] == chunk["unit"] and len(text) >= 100 and chunks[-1]["text"][-100:] == text[:100]:
            text = text[100:]
        if text.strip():
            chunks.append({"unit":chunk["unit"],"text":text})
    return chunks


def _key(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", value)).casefold()


def discover(row):
    chunks = chunks_of(row)
    units = list(dict.fromkeys(c["unit"] for c in chunks))
    entries = []
    path = Path(row["path"])
    if path.suffix.lower() == ".docx" and path.is_file():
        for index, paragraph in enumerate(Document(path).paragraphs):
            match = re.search(r"(?:heading|标题)\s*(\d+)", paragraph.style.name, re.I)
            if match and paragraph.text.strip():
                unit = f"第 {index+1} 段"
                position = next((i for i,c in enumerate(chunks) if c["unit"] == unit), None)
                if position is not None:
                    entries.append({"title":paragraph.text.strip()[:160], "level":int(match[1]), "start":(position,0)})
    if path.suffix.lower() == ".pdf" and path.is_file():
        with fitz.open(path) as document:
            toc = document.get_toc()
        for level, title, page in toc:
            indices = [i for i,c in enumerate(chunks) if c["unit"] == f"第 {page} 页"]
            if not indices:
                continue
            start = (indices[0],0)
            # If two chapters share a physical page, prefer the title's text offset.
            for i in indices:
                found = chunks[i]["text"].find(title.strip())
                if found >= 0:
                    start = (i,found)
                    break
            entries.append({"title":title.strip()[:160], "level":level, "start":start})
    if path.suffix.lower() == ".pptx" and path.is_file():
        for index,slide in enumerate(Presentation(path).slides):
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            if HEADING.fullmatch(title):
                position = next((i for i,c in enumerate(chunks) if c["unit"] == f"第 {index+1} 张幻灯片"),None)
                if position is not None:
                    marker = re.match(r'第[^章节篇部]+([章节篇部])',title)
                    level = {'部':0,'篇':0,'章':1,'节':2}.get(marker[1],1) if marker else 1
                    entries.append({"title":title[:160],"level":level,"start":(position,0)})
    if not entries and not (path.suffix.lower() == ".pptx" and path.is_file()):
        seen = set()
        for index, chunk in enumerate(chunks):
            if re.search(r"目\s*录|contents", '\n'.join(chunk["text"].splitlines()[:6]), re.I):
                continue
            offset = 0
            for line in chunk["text"].splitlines(keepends=True):
                title = line.strip().lstrip('#').strip()
                if HEADING.fullmatch(title) and title not in seen:
                    seen.add(title)
                    marker = re.match(r'第[^章节篇部]+([章节篇部])',title)
                    level = {'部':0,'篇':0,'章':1,'节':2}.get(marker[1],1) if marker else 1
                    entries.append({"title":title[:160], "level":level, "start":(index,offset)})
                offset += len(line)
    entries.sort(key=lambda e:e["start"])
    # Reject unresolved same-position boundaries rather than claiming an exact chapter.
    counts = {}
    for entry in entries:
        counts[entry["start"]] = counts.get(entry["start"],0)+1
    entries = [e for e in entries if counts[e["start"]] == 1]
    for index, entry in enumerate(entries):
        entry["id"] = f"C{index+1}"
        entry["end"] = next((other["start"] for other in entries[index+1:] if other["level"] <= entry["level"]), (len(chunks),0))
        entry["start_unit"] = chunks[entry["start"][0]]["unit"]
        last = entry["end"][0] if entry["end"][1] else entry["end"][0]-1
        entry["end_unit"] = chunks[min(max(last,0),len(chunks)-1)]["unit"]
    return chunks, units, entries


def catalog(row):
    _, units, entries = discover(row)
    return {"material_id":row["id"], "title":row["title"], "units":units,
            "chapters":[{key:e[key] for key in ("id","title","level","start_unit","end_unit")} for e in entries],
            "hint":"按文件目录或明确标题识别；请核对范围。没有可靠标题时可手动选择页码、幻灯片或段落。"}


def select(row, spec):
    chunks, units, entries = discover(row)
    if spec.chapter_id:
        if spec.start_unit or spec.end_unit:
            raise HTTPException(400,"章节与手动位置范围请二选一。")
        entry = next((e for e in entries if e["id"] == spec.chapter_id), None)
        if not entry:
            raise HTTPException(400,"该章节不存在或边界不明确，请重新选择章节或手动范围。")
        result = []
        for index, chunk in enumerate(chunks):
            if index < entry["start"][0] or index > entry["end"][0]:
                continue
            start = entry["start"][1] if index == entry["start"][0] else 0
            end = entry["end"][1] if index == entry["end"][0] else len(chunk["text"])
            if end > start and chunk["text"][start:end].strip():
                result.append({"unit":chunk["unit"],"text":chunk["text"][start:end]})
        label = entry["title"]
        first, last = entry["start_unit"], entry["end_unit"]
    else:
        if not spec.start_unit or not spec.end_unit:
            raise HTTPException(400,"请同时选择范围的开始与结束位置。")
        def resolve(label):
            matches = [u for u in units if _key(u) == _key(label)]
            if len(matches) != 1:
                raise HTTPException(400,"所选位置未提取到文字或不存在，请从可用位置中重新选择。")
            return matches[0]
        first, last = resolve(spec.start_unit), resolve(spec.end_unit)
        if units.index(first) > units.index(last):
            raise HTTPException(400,"范围的开始位置不能在结束位置之后。")
        allowed = set(units[units.index(first):units.index(last)+1])
        result = [c for c in chunks if c["unit"] in allowed]
        label = "自选范围"
    if not result:
        raise HTTPException(400,"所选范围没有可用文字，请检查原文或换一个范围。")
    return result, {"material_id":row["id"],"title":row["title"],"label":label,"start_unit":first,"end_unit":last}
