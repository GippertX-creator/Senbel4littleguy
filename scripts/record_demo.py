"""Record actual browser operations using an isolated, clearly labelled model stub.

Developer-only dependencies: Playwright, Chromium, matplotlib, python-docx,
python-pptx and ffmpeg. The installed workbench does not need these demo tools.
"""
import argparse
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from docx import Document
from playwright.sync_api import sync_playwright, expect
from pptx import Presentation
from pptx.util import Inches, Pt as PptPt

from browser_smoke import free_port

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
EXAMPLES = DOCS / "examples"
IMAGES = DOCS / "images"
MEDIA = DOCS / "media"


def examples():
    """Small, deliberately simplified materials, never a laboratory protocol."""
    EXAMPLES.mkdir(parents=True, exist_ok=True)
    content = [
        ("软件操作练习材料：酶活性", "这是一份为学习工作台录制演示而编写的简化材料，不是完整课程讲义，不可作为真实实验操作方案。"),
        ("酶与催化", "酶能加快生化反应的速率。温度与 pH 均可能影响酶活性。讨论实验时，应区分预期趋势和实际观察。"),
        ("温度与 pH：两个不同变量", "温度描述体系的冷热程度；pH 反映酸碱程度。二者不能互相替代。研究温度影响时，应控制 pH 等其他条件；研究 pH 时，应控制温度等其他条件。"),
        ("对照与控制变量", "比较实验组和对照组时，应明确研究变量，并尽可能保持其他条件一致。同一次比较中同时改变多个条件，通常难以区分各因素的作用。"),
        ("实验报告与真实记录", "实验数据、现象、结果和结论应来自真实实验。没有完成实验时，不填写虚构数值，也不把预期现象当作观察结果。报告固定内容可以根据教师提供的材料事先整理。"),
    ]
    ppt = Presentation()
    ppt.slide_width = Inches(12)
    ppt.slide_height = Inches(6.75)
    for title, body in content:
        slide = ppt.slides.add_slide(ppt.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Microsoft YaHei"
                        run.font.size = PptPt(28 if shape == slide.shapes.title else 22)
    ppt.save(EXAMPLES / "01-演示课程材料.pptx")
    doc = Document()
    doc.add_heading("温度对酶活性影响：软件演示材料", 0)
    doc.add_paragraph("仅演示材料上传和报告生成。不是可直接开展实验的完整方案；试剂用量、温度范围、时间及操作要求应以教师的正式实验指导书为准。")
    for title, body in [
        ("实验目的", "理解控制变量的思想，讨论温度与酶活性的关系。"),
        ("实验原理", "温度与 pH 均可能影响酶活性。比较温度的影响时，需要控制其他条件。"),
        ("仪器与试剂", "材料未列出完整清单，请补充教师正式指导书。"),
        ("实验步骤", "仅说明设计原则：明确温度这一研究变量，设置可比较的组别，尽可能保持 pH 等其他条件一致，并按教师要求记录真实结果。具体操作步骤未提供，请补充。"),
        ("注意事项", "不得把预期趋势写成观察结果。数据、现象和结论必须根据实际实验填写。"),
    ]:
        doc.add_heading(title, 1)
        doc.add_paragraph(body)
    doc.save(EXAMPLES / "02-演示实验材料.docx")
    doc = Document()
    doc.add_heading("{{实验名称}}", 0)
    doc.add_paragraph("姓名：________    学号：________    日期：________")
    for field in ["实验目的", "实验原理", "仪器与试剂", "实验步骤", "注意事项", "实验数据", "实验结果", "分析与讨论", "实验结论"]:
        doc.add_heading(field, 1)
        doc.add_paragraph("{{" + field + "}}")
    doc.save(EXAMPLES / "03-演示报告模板.docx")
    # A new schematic drawn with standard plotting tools, not a micrograph.
    os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "data" / "demo-matplotlib"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Ellipse
    fig = plt.figure(figsize=(8, 5), dpi=120, facecolor="#fffaf7")
    ax = fig.add_axes((0, 0, 1, 1)); ax.set(xlim=(0, 1), ylim=(0, 1)); ax.axis("off")
    ax.add_patch(Ellipse((.5, .5), .86, .83, facecolor="#f4eafa", edgecolor="#ac89bc", linewidth=6))
    ax.add_patch(Ellipse((.48, .52), .25, .30, facecolor="#b89aca", edgecolor="#8d6b9f", linewidth=3))
    ax.add_patch(Ellipse((.5, .53), .065, .09, facecolor="#80618f", edgecolor="none"))
    for x, y, angle in [(.27, .69, 25), (.71, .65, -30), (.7, .31, 15), (.3, .3, -20)]:
        ax.add_patch(Ellipse((x, y), .12, .075, angle=angle, facecolor="#ecc5b0", edgecolor="#c89c83", linewidth=2))
    fig.savefig(EXAMPLES / "04-演示细胞示意图.png", dpi=120)
    plt.close(fig)
    (EXAMPLES / "请先阅读.txt").write_text("这些文件仅供软件操作练习，不替代教师课件、实验指导书或标准图谱。细胞图片是简化示意图，不是显微照片。视频中固定示例由本机模拟服务返回，不代表真实模型效果。\n", encoding="utf-8")


class DemoModel(BaseHTTPRequestHandler):
    """Only the external API is simulated. Everything in the app runs normally."""
    def log_message(self, *args):
        pass

    def do_POST(self):
        payload = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        prompt = payload["messages"][-1]["content"]
        time.sleep(1.5)
        matches = list(dict.fromkeys(re.findall(r"\[材料 ([0-9a-f]{32}) \| [^\n]+? \| ([^\]]+)\]", prompt)))
        refs = [{"material_id": m[0], "unit": m[1]} for m in matches]
        source = refs[:1]
        def cited(**kwargs):
            return {**kwargs, "sources": source}
        bank = [
            cited(stem="研究温度对酶活性的影响时，pH 应如何处理？", options=["尽可能保持一致", "和温度同时随意改变", "无需记录", "只能改变 pH"], answer=0, explanation="为了比较温度的作用，应尽可能控制 pH 等其他条件。", topic="控制变量"),
            cited(stem="下列哪一项属于实验后才能填写的内容？", options=["教师提供的实验目的", "材料中已写明的原理", "本次实验的真实观察结果", "指导书中的注意事项"], answer=2, explanation="实际观察结果必须来自本次实验，不能由预期趋势代替。", topic="真实记录"),
            cited(stem="研究 pH 的作用时，下列哪项设计更便于解释结果？", options=["同时改变温度与 pH", "只比较不同 pH，并控制温度等条件", "不设可比较的组别", "不记录条件"], answer=1, explanation="一次比较围绕一个研究变量展开，更便于分析该变量的影响。", topic="实验设计"),
            cited(stem="温度和 pH 的关系，哪项描述符合材料？", options=["二者是同一个变量", "pH 描述冷热程度", "温度描述酸碱程度", "二者是不同条件，均可能影响酶活性"], answer=3, explanation="材料区分了温度与 pH，并说明两者均可能影响酶活性。", topic="易混概念"),
            cited(stem="尚未完成实验时，报告中的实验结论应该怎样处理？", options=["根据预期趋势写成真实结论", "先编一组合理数据", "保留待填，完成实验后据实填写", "复制别人的结果"], answer=2, explanation="数据和结论必须有真实实验依据。", topic="报告填写"),
        ]
        if '"title": "Outline"' in prompt:
            result = {"summary": "围绕酶活性、控制变量与真实记录，先理解机制，再检查实验设计。", "items": [
                cited(term="酶与催化", explanation="酶能加快生化反应速率。温度和 pH 均可能影响酶活性。", importance="核心", recall_question="哪些条件可能影响酶活性？"),
                cited(term="研究变量与控制变量", explanation="研究温度影响时，尽可能保持 pH 等其他条件一致。", importance="核心", recall_question="为什么不宜同时任意改变温度与 pH？"),
                cited(term="真实实验记录", explanation="预期趋势与实际观察不同，实验结果和结论应来自真实实验。", importance="核心", recall_question="没有完成实验时，哪些报告内容应该留空？"),
            ], "connections": [cited(chain="明确研究变量 → 控制其他条件 → 比较真实记录", explanation="通过可比较的实验设计，减少对结果解释的混淆。") ]}
        elif '"title": "Quiz"' in prompt:
            count = int(re.search(r"生成 (\d+) 道", prompt)[1]); result = {"questions": [bank[i % len(bank)] for i in range(count)]}
        elif '"title": "Comparison"' in prompt:
            result = {"left_name": "温度", "right_name": "pH", "rows": [cited(dimension="描述什么", left="体系的冷热程度", right="体系的酸碱程度"), cited(dimension="作为研究变量时", left="比较温度，控制 pH 等其他条件", right="比较 pH，控制温度等其他条件")], "memory_tip": "先说清本次研究哪个变量，再列出需要保持一致的条件。", "practice": [bank[0], bank[2]]}
        elif '"title": "Feedback"' in prompt:
            result = {"understood": ["你提到了温度和 pH 均可能影响酶活性。"], "corrections": [cited(student_claim="为了比较温度的影响，可以同时随意改变 pH。", feedback="同时任意改变多个条件，通常难以区分它们各自的作用。研究温度时，应控制 pH 等其他条件。")], "missing": [], "next_question": "如果温度是本次的研究变量，你准备让哪些条件保持一致？"}
        elif '"title": "Report"' in prompt:
            result = {"title": "温度对酶活性的影响", "sections": [cited(key=k, title=t, content=c) for k,t,c in [
                ("purpose", "实验目的", "理解控制变量的思想，讨论温度与酶活性的关系。"), ("principle", "实验原理", "温度与 pH 均可能影响酶活性。比较温度时，需要控制其他条件。"),
                ("equipment", "仪器与试剂", "材料未列出完整清单，请补充教师正式指导书。"), ("steps", "实验步骤", "明确温度这一研究变量，设置可比较的组别，并尽可能保持 pH 等其他条件一致。具体用量、温度范围和操作步骤未提供，请补充。"),
                ("precautions", "注意事项", "不得把预期趋势写成实际观察。数据、结果与结论根据真实实验填写。")]], "pre_questions": bank[:3]}
        elif '"title": "Answer"' in prompt:
            result = cited(answer="因为温度与 pH 都可能影响酶活性。如果两者同时任意变化，就难以区分观察到的差异来自哪个条件。研究温度时，应尽可能保持 pH 等其他条件一致，并记录真实结果。", follow_up="如果你改为研究 pH，应该怎样处理温度？")
        else:
            result = "OK"
        content = result if isinstance(result, str) else json.dumps(result, ensure_ascii=False)
        body = json.dumps({"choices": [{"finish_reason": "stop", "message": {"content": content}}]}, ensure_ascii=False).encode()
        self.send_response(200); self.send_header("Content-Type", "application/json"); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)


OVERLAY = """() => {
const style = document.createElement('style');
style.textContent = `#demo-guide{position:fixed;bottom:0;left:0;right:0;height:96px;padding:15px 36px;background:#443652f5;color:#fff;z-index:10000;box-sizing:border-box;font-family:'Noto Sans CJK SC','Microsoft YaHei',sans-serif;pointer-events:none;border-top:3px solid #d8b9e8}#demo-guide strong{display:block;font-size:19px;margin-bottom:7px}#demo-guide p{font-size:14px;line-height:1.6;margin:0;color:#efdfef}#demo-guide small{position:absolute;right:28px;top:16px;font-size:11px;color:#ddc8df}#demo-badge{position:fixed;right:26px;top:83px;background:#fff5e9f0;border:1px solid #e6c7a0;border-radius:8px;padding:6px 10px;font:11px 'Noto Sans CJK SC',sans-serif;color:#996d3e;z-index:9999;pointer-events:none}#demo-cursor{position:fixed;width:32px;height:32px;border:2px solid #e4ab67;background:#f3bc6b22;border-radius:50%;left:-100px;top:-100px;transform:translate(-50%,-50%);pointer-events:none;z-index:10001;transition:width .12s,height .12s}.cat-companion{bottom:108px!important}#toast{bottom:110px!important}html{scroll-padding-bottom:120px}`;
document.head.appendChild(style);
const guide=document.createElement('div');guide.id='demo-guide';guide.innerHTML='<strong></strong><p></p><small>小家伙的小家伙 · 中文操作演示</small>';document.body.appendChild(guide);
const badge=document.createElement('div');badge.id='demo-badge';badge.textContent='模拟 API 演示 · 固定示例，非真实模型效果';document.body.appendChild(badge);
const cursor=document.createElement('div');cursor.id='demo-cursor';document.body.appendChild(cursor);
document.addEventListener('mousemove',e=>{cursor.style.left=e.clientX+'px';cursor.style.top=e.clientY+'px'});
document.addEventListener('mousedown',()=>{cursor.style.width='22px';cursor.style.height='22px'});
document.addEventListener('mouseup',()=>{cursor.style.width='32px';cursor.style.height='32px'});
} """


class Director:
    def __init__(self, page, speed):
        self.page = page; self.speed = speed; self.start = time.monotonic(); self.chapters = []

    def pause(self, seconds=3):
        self.page.wait_for_timeout(seconds * 1000 * self.speed)

    def chapter(self, title, text):
        self.chapters.append({"start": round(time.monotonic() - self.start, 2), "title": title, "text": text})
        self.page.evaluate("([t,s])=>{document.querySelector('#demo-guide strong').textContent=t;document.querySelector('#demo-guide p').textContent=s}", [title, text])
        self.pause(3)

    def click(self, selector):
        node = self.page.locator(selector)
        node.evaluate("node => node.scrollIntoView({block:'center',behavior:'instant'})"); node.hover(); self.pause(.55); node.click(); self.pause(.75)

    def fill(self, selector, text):
        node = self.page.locator(selector);node.evaluate("node => node.scrollIntoView({block:'center',behavior:'instant'})");node.hover();node.fill(text);self.pause(1)

    def nav(self, view):
        self.click(f'#navigation a[href="#{view}"]'); self.page.evaluate('window.scrollTo(0,0)')

    def shot(self, name):
        # Documentation screenshots omit recording annotations and remain actual UI captures.
        self.page.evaluate("document.querySelectorAll('#demo-guide,#demo-badge,#demo-cursor').forEach(e=>e.style.visibility='hidden')")
        self.page.screenshot(path=str(IMAGES / name), animations="disabled")
        self.page.evaluate("document.querySelectorAll('#demo-guide,#demo-badge,#demo-cursor').forEach(e=>e.style.visibility='')")


def stamp(seconds):
    ms = int(seconds * 1000);return f"{ms//3600000:02}:{ms//60000%60:02}:{ms//1000%60:02},{ms%1000:03}"


def run(speed=1.0):
    for directory in (IMAGES, MEDIA): directory.mkdir(parents=True, exist_ok=True)
    examples()
    model = ThreadingHTTPServer(("127.0.0.1", 0), DemoModel)
    threading.Thread(target=model.serve_forever, daemon=True).start()
    errors = []
    with tempfile.TemporaryDirectory() as tmp:
        port=free_port();url=f"http://127.0.0.1:{port}"
        python=ROOT/".venv"/("Scripts/python.exe" if os.name=="nt" else "bin/python")
        if not python.exists(): python=Path(sys.executable)
        env={**os.environ, "MEDSTUDY_DATA": str(Path(tmp)/"data")}
        with (Path(tmp)/"server.log").open("w") as log:
            process=subprocess.Popen([str(python),"-m","uvicorn","app.main:app","--host","127.0.0.1","--port",str(port)],cwd=ROOT,env=env,stdout=log,stderr=log)
            try:
                for _ in range(80):
                    if process.poll() is not None: raise RuntimeError((Path(tmp)/"server.log").read_text())
                    try:
                        with urllib.request.urlopen(url+"/api/health",timeout=1): break
                    except OSError: time.sleep(.15)
                with sync_playwright() as p:
                    browser=p.chromium.launch(headless=True,executable_path=shutil.which("chromium"),args=["--no-sandbox"])
                    context=browser.new_context(viewport={"width":1440,"height":960},device_scale_factor=1,record_video_dir=str(Path(tmp)/"recording"),record_video_size={"width":1440,"height":960},accept_downloads=True)
                    page=context.new_page();page.on("pageerror",lambda err:errors.append(str(err)))
                    page.goto(url);expect(page.locator("h1")).to_have_text("小家伙的小家伙")
                    page.wait_for_function("document.querySelector('.hero-cat img').complete")
                    page.evaluate(OVERLAY);d=Director(page,speed)
                    d.chapter("01 / 欢迎来到你的学习工作台", "这是网页操作实录。Windows 安装步骤请看说明书；本片使用模拟 API，不展示真实模型效果。")
                    d.pause(5); d.shot("01-home.png")
                    d.chapter("02 / 小猫陪伴与大字模式", "右上角“大字”可放大文字；小猫可以摸摸，也能暂停动画。它不会反复弹出提示遮挡内容。")
                    d.click(".hero-cat-button"); d.pause(3);d.click("#cat-motion-toggle");d.pause(1.5);d.click("#cat-motion-toggle");d.click("[data-action=close-cat]");d.click("#text-size-toggle");d.pause(2);d.click("#text-size-toggle")
                    d.chapter("03 / 第一次使用：填写 API 三项信息", "基础地址、模型名称、API 密钥来自你的服务商。不要照抄片中的本机演示地址和模型名。")
                    d.nav("settings");d.fill("[name=base_url]",f"http://127.0.0.1:{model.server_port}/v1");d.fill("[name=model]","demo-model-not-a-real-provider");d.fill("[name=api_key]","DEMO-ONLY-NOT-A-REAL-KEY")
                    d.shot("02-api-settings.png");d.click("[data-action=test-api]");expect(page.locator("#toast")).to_contain_text("API 请求成功");d.pause(4)
                    d.chapter("04 / 先选课程，再上传材料", "这里选择生物化学。课程材料、实验材料、报告模板用途不同，请按用途分别上传。")
                    page.locator("#course-select").select_option("biochemistry");d.pause(1);d.nav("materials")
                    for kind, filename in [("course","01-演示课程材料.pptx"),("experiment","02-演示实验材料.docx"),("template","03-演示报告模板.docx")]:
                        page.locator("[name=kind]").select_option(kind)
                        page.locator("[name=files]").set_input_files(EXAMPLES/filename);d.pause(1)
                        d.click("#upload-form button[type=submit]");expect(page.locator("#toast")).to_contain_text("已解析");d.pause(1)
                    page.evaluate("window.scrollTo(0,document.body.scrollHeight)");d.pause(3);d.shot("03-materials.png")
                    d.chapter("05 / 知识清单：先看结构，再主动回忆", "选择材料和学习范围，生成后可点击来源核对原文，再把问题加入复习。")
                    d.nav("outline");d.fill("#outline-form [name=topic]","酶活性与控制变量");d.click("#outline-form button[type=submit]");expect(page.locator(".knowledge-card")).to_have_count(3);d.pause(4);d.shot("04-outline.png")
                    d.click(".knowledge-card .source-link >> nth=0");expect(page.locator("#source-dialog")).to_be_visible();d.pause(3);d.shot("05-source.png");d.click("#close-dialog");d.click("[data-action=recall]");d.pause(2)
                    d.chapter("06 / 课程助教：从一个为什么开始", "输入你没弄懂的问题。小猫思考后给出解释与来源；看完答案，还可以接着追问。")
                    d.nav("chat");d.fill("[name=message]","为什么研究温度对酶活性的影响时，要保持 pH 等其他条件一致？");d.click("#chat-form button[type=submit]");expect(page.locator(".message.assistant")).to_contain_text("温度与 pH");d.pause(4);d.shot("06-tutor.png")
                    d.chapter("07 / 讲给助教听：找出理解中的漏洞", "先用自己的话解释。反馈会区分已经讲清楚和需要修正的部分，你可以修改后再提交。")
                    d.nav("teach");d.fill("[name=topic]","酶活性实验的控制变量");d.fill("[name=explanation]","温度和 pH 都可能影响酶活性。为了比较温度的影响，可以同时随意改变 pH。");d.click("#teach-form button[type=submit]");expect(page.locator(".result-container")).to_contain_text("研究温度时");d.pause(4);d.shot("07-teach.png")
                    d.chapter("08 / 实验报告：生成固定内容，保留真实记录", "选择报告模板并生成。数据、观察、结果、讨论和结论仍需你根据实际实验填写。")
                    d.nav("report");d.fill("[name=title]","温度对酶活性的影响")
                    page.locator("[name=template_id]").select_option(label="03-演示报告模板.docx");d.pause(1);d.shot("08-report-form.png")
                    d.click("#report-form button[type=submit]");expect(page.locator(".report-section")).to_have_count(5);d.pause(4);d.shot("09-report-result.png")
                    with page.expect_download() as download:d.click('.result-container a[href^="/api/exports/"]')
                    download.value.save_as(str(MEDIA/"演示生成的实验报告.docx"));d.pause(2)
                    d.chapter("09 / 实验学习卡：作答后才显示解析", "报告关联原理小测。请答完每一道题，再点击提交；答错的题会进入今日复习。")
                    d.click(".result-container [data-action=artifact]");expect(page.locator(".quiz-card")).to_have_count(3)
                    for index,answer in enumerate([1,2,1]):d.click(f'[name=q{index}][value="{answer}"]')
                    d.click("#grade-form button[type=submit]");expect(page.locator(".score-strip strong")).to_have_text("2 / 3");d.pause(4);d.shot("10-quiz.png")
                    d.chapter("10 / 易混知识：先辨别，再对照", "写下你对两个概念的初步判断，再看比较表，最后完成新的情境题。")
                    d.nav("practice");d.click("[data-tab=compare]");d.fill("[name=left]","温度");d.fill("[name=right]","pH");d.fill("[name=guess]","它们是两个不同条件，研究其中一个时需要控制另一个。");d.click("#compare-form button[type=submit]");expect(page.locator(".compare-table")).to_be_visible();d.pause(4);d.shot("11-compare.png")
                    d.chapter("11 / 指图学习：先建立可靠标注", "上传无答案标签的图片，点击结构填写名称。演示使用简化示意图，不是显微照片。")
                    d.nav("figures");page.locator('#figure-upload-form [name=file]').set_input_files(EXAMPLES/"04-演示细胞示意图.png");d.pause(1);d.click("#figure-upload-form button[type=submit]");expect(page.locator("#figure-stage img")).to_be_visible()
                    stage=page.locator("#figure-stage");stage.scroll_into_view_if_needed();d.pause(1);box=stage.bounding_box();page.mouse.click(box["x"]+box["width"]*.48,box["y"]+box["height"]*.48);d.pause(1)
                    d.fill('[name^="name-"]',"细胞核 / nucleus");d.fill('[name^="explanation-"]',"演示示意图中心偏左的紫色结构。正式学习请使用教师图谱确认。");page.locator('[name=confirmed]').check();d.click("#labels-form button[type=submit]");expect(page.locator('#toast')).to_contain_text('标注已保存');d.shot("12-figure-labels.png")
                    d.chapter("12 / 隐藏标签：辨认结构并核对", "系统只隐藏你添加的标签，不会擦除原图的文字。评分按已确认名称及别名匹配。")
                    d.click('[data-mode=practice]');d.fill('#figure-grade-form input',"nucleus");d.click('#figure-grade-form button[type=submit]');expect(page.locator('#figure-grade-form')).to_contain_text('1 / 1');d.pause(4);d.shot('13-figure-practice.png')
                    d.chapter("13 / 今日复习：先回忆，再看参考", "先写下你记得的内容，再查看参考。能回忆就安排下一次，仍不熟悉就明天再练。")
                    d.nav('review');d.fill('[name=answer]',"研究温度时控制其他条件，实验结果必须真实记录。");d.click('#review-form button[type=submit]');expect(page.locator('#review-form')).to_contain_text('参考内容');d.pause(3);d.shot('14-review.png');d.click('[data-remembered=true]');expect(page.locator('#toast')).to_contain_text('下次复习');d.pause(2)
                    d.chapter("14 / 下次怎么打开？资料怎么保留？", "每天双击 start.bat，保持启动窗口打开；服务重启后重填密钥。备份前关闭服务，再复制整个 data 文件夹。")
                    d.nav('home');d.pause(5)
                    d.chapter("15 / 现在可以自己试一遍了", "先用 docs/examples 里的练习材料操作。真实 API、正式课件和医学内容，请以你实际的服务与课程要求为准。")
                    d.pause(6)
                    duration=time.monotonic()-d.start
                    recording=page.video;context.close();raw=MEDIA/"demo-original.webm";recording.save_as(str(raw));browser.close()
                    if errors:raise RuntimeError("Browser errors: "+str(errors))
                    timeline={"duration":round(duration,2),"chapters":d.chapters,"disclosure":"真实网页操作；模型 API 为本机固定示例。无真实密钥，无真实模型效果验证。Windows 安装未录屏。"}
                    (MEDIA/"演示章节.json").write_text(json.dumps(timeline,ensure_ascii=False,indent=2),encoding="utf-8")
            finally:
                process.terminate()
                try:process.wait(timeout=10)
                except subprocess.TimeoutExpired:process.kill();process.wait()
    model.shutdown();model.server_close()
    metadata=[";FFMETADATA1", "title=小家伙的小家伙：零基础操作演示", "comment=真实网页录屏，使用固定模拟API返回；非真实模型效果。中文字幕，无配音。"]
    subtitles=[]
    for index, chapter in enumerate(timeline["chapters"]):
        end=timeline["chapters"][index+1]["start"] if index+1<len(timeline["chapters"]) else timeline["duration"]
        metadata.extend(["[CHAPTER]","TIMEBASE=1/1000",f"START={int(chapter['start']*1000)}",f"END={int(end*1000)}","title="+chapter["title"]])
        subtitles.append(f"{index+1}\n{stamp(chapter['start'])} --> {stamp(end)}\n{chapter['title']}\n{chapter['text']}\n")
    (MEDIA/"演示字幕.srt").write_text("\n".join(subtitles),encoding="utf-8")
    meta=MEDIA/"chapters.ffmetadata";meta.write_text("\n".join(metadata),encoding="utf-8")
    output=MEDIA/"小家伙的小家伙-操作演示.mp4"
    subprocess.run(["ffmpeg","-hide_banner","-loglevel","warning","-y","-i",str(MEDIA/"demo-original.webm"),"-i",str(meta),"-map_metadata","1","-map_chapters","1","-an","-c:v","libx264","-preset","fast","-crf","24","-pix_fmt","yuv420p","-r","25","-movflags","+faststart",str(output)],check=True)
    (MEDIA/"demo-original.webm").unlink();meta.unlink()
    print("Demo complete:",output,"duration:",timeline["duration"],"seconds; screenshots:",len(list(IMAGES.glob('[0-9][0-9]-*.png'))))


if __name__=="__main__":
    parser=argparse.ArgumentParser();parser.add_argument('--speed',type=float,default=1.0,help='Pause multiplier; 1 is the normal teaching pace.')
    args=parser.parse_args();run(args.speed)
