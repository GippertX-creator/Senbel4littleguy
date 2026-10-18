"""Browser verification of chapter scope, Chinese PDF and the learning agent.

Uses a local synthetic model; it does not validate any real provider's reasoning.
"""
import argparse
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import threading
import time
import urllib.request
from http.server import ThreadingHTTPServer

import fitz
from docx import Document
from playwright.sync_api import expect, sync_playwright

from browser_smoke import free_port, ModelStub

ROOT = Path(__file__).resolve().parent.parent


class AgentStub(ModelStub):
    def do_POST(self):
        raw = self.rfile.read(int(self.headers['Content-Length']))
        payload = json.loads(raw)
        self.server.requests.append(payload)
        prompt = payload['messages'][-1]['content']
        result = None
        if '"title": "AgentPlan"' in prompt:
            message = prompt.split('学生当前要求：',1)[1].split('\n输出 JSON schema',1)[0]
            if message == '帮我复习':
                result = {'status':'clarify','message':'你想先做知识清单，还是通过小测检查理解？','tasks':[]}
            else:
                tasks = [{'action':'quiz','topic':'受体与第二信使','count':1}] if '再出1' in message else [
                    {'action':'outline','material_key':'M1','chapter_key':'C2','topic':'受体与第二信使','instructions':'列两个重点'},
                    {'action':'quiz','material_key':'M1','chapter_key':'C2','topic':'受体与第二信使','count':2}]
                result = {'status':'execute','message':'我会只用第二章，按你的要求整理并提供练习。','tasks':tasks}
        elif '"title": "Outline"' in prompt or '"title": "Quiz"' in prompt:
            source_ids = re.findall(r'^\[来源 (S\d+)\]$',prompt,re.M)
            sources = [{'source_id':source_ids[min(1,len(source_ids)-1)]}]
            if '"title": "Outline"' in prompt:
                result = {'summary':'本次只整理第二章的受体与第二信使。','items':[
                    {'term':'受体','explanation':'受体识别信号并启动细胞内传递。','importance':'核心','recall_question':'信号首先由谁识别？','sources':sources},
                    {'term':'第二信使','explanation':'第二信使在细胞内传递和放大信号。','importance':'核心','recall_question':'第二信使起什么作用？','sources':sources}],
                    'connections':[{'chain':'信号 → 受体 → 胞内传递','explanation':'受体将信号与细胞内反应联系起来。','sources':sources}]}
            else:
                count = int(re.search(r'生成 (\d+) 道',prompt)[1])
                question = {'stem':'受体在信号转导中的作用是什么？','options':['识别信号并启动胞内传递','代替全部第二信使','生成任意实验数据','与信号传递无关'],'answer':0,'explanation':'材料说明受体识别信号并启动细胞内传递。','topic':'受体','sources':sources}
                result = {'questions':[question for _ in range(count)]}
        if result is None:
            self.rfile = io.BytesIO(raw)
            return super().do_POST()
        body = json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(result,ensure_ascii=False)}}]},ensure_ascii=False).encode()
        self.send_response(200);self.send_header('Content-Type','application/json');self.send_header('Content-Length',str(len(body)));self.end_headers();self.wfile.write(body)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,default=Path('/tmp/flexible-learning'))
    args = parser.parse_args();args.output.mkdir(parents=True,exist_ok=True)
    model = ThreadingHTTPServer(('127.0.0.1',0),AgentStub)
    model.requests=[];model.truncated_probes=0;model.citation_probes=0
    threading.Thread(target=model.serve_forever,daemon=True).start()
    errors=[]
    with tempfile.TemporaryDirectory() as tmp:
        document=Document()
        for title,text in [('第一章 细胞结构','FIRST_CHAPTER_ONLY 第一章介绍细胞基本结构。'),
                           ('第二章 信号转导','SECOND_CHAPTER_ONLY 受体识别信号并启动细胞内传递。第二信使在细胞内传递和放大信号。'),
                           ('第三章 细胞周期','THIRD_CHAPTER_ONLY 第三章介绍细胞周期。')]:
            document.add_heading(title,1);document.add_paragraph(text)
        file=Path(tmp)/'演示课程.docx';document.save(file)
        port=free_port();url=f'http://127.0.0.1:{port}'
        with (Path(tmp)/'server.log').open('w') as log:
            server=subprocess.Popen([str(ROOT/'.venv/bin/python'),'-m','uvicorn','app.main:app','--host','127.0.0.1','--port',str(port)],cwd=ROOT,env={**os.environ,'MEDSTUDY_DATA':str(Path(tmp)/'data')},stdout=log,stderr=log)
            try:
                for _ in range(60):
                    if server.poll() is not None:raise RuntimeError((Path(tmp)/'server.log').read_text())
                    try:urllib.request.urlopen(url+'/api/health',timeout=1).close();break
                    except OSError:time.sleep(.2)
                with sync_playwright() as p:
                    browser=p.chromium.launch(headless=True,executable_path=shutil.which('chromium'),args=['--no-sandbox'])
                    page=browser.new_page(viewport={'width':1440,'height':1050},accept_downloads=True)
                    page.on('pageerror',lambda error:errors.append(str(error)))
                    page.goto(url)
                    def nav(view):
                        page.locator(f'#navigation a[href="#{view}"]').click()
                    nav('settings')
                    page.locator('[name=base_url]').fill(f'http://127.0.0.1:{model.server_port}/v1')
                    page.locator('[name=model]').fill('local-synthetic-model')
                    page.locator('[data-action=test-api]').click()
                    expect(page.locator('#toast')).to_contain_text('API 请求成功')
                    page.locator('#course-select').select_option('cell-biology')
                    nav('materials')
                    page.locator('[name=files]').set_input_files(file)
                    page.locator('#upload-form button[type=submit]').click()
                    expect(page.locator('.list-info h3')).to_have_text(file.name)
                    nav('outline')
                    page.locator('.range-selector summary').click()
                    page.locator('#scope-mode').select_option('chapter')
                    expect(page.locator('#scope-chapter option')).to_have_count(4)
                    page.locator('#scope-chapter').select_option('C2')
                    page.locator('[name=topic]').fill('受体与第二信使')
                    page.locator('[name=instructions]').fill('只列两个重点，突出因果链。')
                    page.locator('#outline-form button[type=submit]').click()
                    expect(page.locator('.knowledge-card')).to_have_count(2)
                    expect(page.locator('.scope-coverage')).to_contain_text('第二章 信号转导')
                    page.evaluate("window.scrollTo({top:0,behavior:'instant'})")
                    page.wait_for_timeout(250)
                    page.screenshot(path=str(args.output/'chapter-outline.png'),full_page=True,animations='disabled')
                    with page.expect_download() as download:page.locator('.pdf-download').click()
                    pdf=Path(tmp)/'outline.pdf';download.value.save_as(pdf)
                    with fitz.open(pdf) as doc:text=''.join(p.get_text() for p in doc)
                    assert '第二章 信号转导' in text and '主动回忆' in text
                    nav('agent')
                    page.locator('#agent-form [name=message]').fill('帮我复习')
                    page.locator('#agent-form button[type=submit]').click()
                    expect(page.locator('.agent-message.assistant')).to_contain_text('知识清单')
                    expect(page.locator('#scope-chapter')).to_have_value('C2')
                    thread_id=page.locator('#agent-thread-select').input_value()
                    assert thread_id
                    page.locator('#agent-form [name=message]').fill('只用第二章，做清单和2道题，并提供 PDF')
                    page.locator('#agent-form button[type=submit]').click()
                    expect(page.locator('.agent-result-card')).to_have_count(2)
                    expect(page.locator('.knowledge-card')).to_have_count(2)
                    page.locator('.agent-result-card [data-action=agent-result]').nth(1).click()
                    expect(page.locator('.quiz-card')).to_have_count(2)
                    for index in range(2):page.locator(f'[name=q{index}][value="0"]').check()
                    page.locator('#grade-form button[type=submit]').click()
                    expect(page.locator('.score-strip')).to_contain_text('2 / 2')
                    page.evaluate('window.scrollTo(0,0)')
                    page.screenshot(path=str(args.output/'agent-workflow.png'),full_page=True,animations='disabled')
                    page.reload()
                    page.locator('#agent-thread-select').select_option(thread_id)
                    expect(page.locator('.agent-message')).to_have_count(4)
                    expect(page.locator('#scope-chapter')).to_have_value('C2')
                    page.locator('#agent-form [name=message]').fill('再出1道题')
                    page.locator('#agent-form button[type=submit]').click()
                    expect(page.locator('.quiz-card')).to_have_count(1)
                    before=len(model.requests)
                    with page.expect_download() as download:page.locator('.agent-result-card a').last.click()
                    download.value.save_as(Path(tmp)/'agent-quiz.pdf')
                    assert len(model.requests)==before,'PDF export must not call the model'
                    page.set_viewport_size({'width':390,'height':844})
                    page.evaluate('window.scrollTo(0,0)')
                    assert page.evaluate('document.documentElement.scrollWidth<=innerWidth'),'Mobile agent overflows'
                    page.screenshot(path=str(args.output/'agent-mobile.png'),full_page=True,animations='disabled')
                    page.locator('#course-select').select_option('pathology')
                    expect(page.locator('.agent-message')).to_have_count(0)
                    expect(page.locator('#agent-thread-select option')).to_have_count(1)
                    browser.close()
                for request in model.requests:
                    prompt=request['messages'][-1]['content']
                    if '"title": "Outline"' in prompt or '"title": "Quiz"' in prompt:
                        assert 'SECOND_CHAPTER_ONLY' in prompt and 'FIRST_CHAPTER_ONLY' not in prompt and 'THIRD_CHAPTER_ONLY' not in prompt
                assert not errors,errors
                print('Flexible learning PASS: exact chapter UI, Chinese PDF, clarification, two-tool execution, quiz grading, scoped follow-up, restored conversation, course isolation, 390px layout; no JavaScript errors.')
            finally:
                server.terminate()
                try:server.wait(timeout=10)
                except subprocess.TimeoutExpired:server.kill();server.wait()
    model.shutdown();model.server_close()


if __name__=='__main__':main()
