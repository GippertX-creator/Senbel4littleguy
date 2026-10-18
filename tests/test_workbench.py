import io
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

import fitz
import httpx
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from app import ai, db
from app.main import app

REAL_ASYNC_CLIENT = httpx.AsyncClient


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path_patch = patch.object(db, "DATA", Path(self.temp.name))
        self.path_patch.start()
        ai.API_KEY = ""
        self.client = TestClient(app)
        self.client.__enter__()
        self.response = None
        self.completion_body = None
        self.completion_queue = []
        self.status = 200
        self.captured = []
        self.transport = httpx.MockTransport(self.upstream)
        self.http_patch = patch("app.ai.httpx.AsyncClient", side_effect=lambda **kwargs: REAL_ASYNC_CLIENT(transport=self.transport, **kwargs))
        self.http_patch.start()
        self.client.put("/api/settings", json={"base_url":"https://example.com/v1", "model":"test-model", "api_key":"test-only-key"})
        result = self.upload("实验.txt", "酶活性受温度和 pH 影响。实验目的：观察温度对酶活性的影响。实验原理：酶催化反应。仪器与试剂：试管与淀粉酶。实验步骤：设置温度梯度，保持其他条件一致。注意事项：使用对照组。".encode())
        self.material = result.json()["id"]
        self.ref = {"material_id":self.material,"unit":"段落 1"}

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.http_patch.stop()
        self.path_patch.stop()
        ai.API_KEY = ""
        self.temp.cleanup()

    def upstream(self, request):
        self.captured.append((request.url, request.headers, json.loads(request.content)))
        if self.completion_queue:
            return httpx.Response(self.status, json=self.completion_queue.pop(0))
        if self.completion_body is not None:
            return httpx.Response(self.status, json=self.completion_body)
        return httpx.Response(self.status, json={"choices":[{"finish_reason":"stop", "message":{"content":self.response if isinstance(self.response,str) else json.dumps(self.response,ensure_ascii=False)}}]})

    def upload(self, name, data, kind="course", course="biochemistry"):
        return self.client.post("/api/materials", data={"course":course,"kind":kind}, files={"file":(name,data)})

    def question(self, answer=0):
        return {"stem":"设计温度实验时，哪些条件应保持一致？", "options":["其他条件","只改变 pH","所有条件同时变化","不设置对照"], "answer":answer,
                "explanation":"保持其他条件一致，才能比较温度的作用。", "topic":"实验对照", "sources":[self.ref.copy()]}

    def request(self, extra=None):
        return {"course":"biochemistry", "material_ids":[self.material], **(extra or {})}

    def completion(self, content, finish="stop"):
        return {"choices":[{"finish_reason":finish,"message":{"content":content,
                "reasoning_content":"private reasoning"}}]}

    def outline_response(self, sources):
        return {'summary':'酶活性学习','items':[{'term':'控制变量','explanation':'保持其他条件一致',
                'importance':'核心','recall_question':'为什么设置对照？','sources':sources}], 'connections':[]}

    def create_quiz(self):
        self.response = {"questions":[self.question(),self.question()]}
        result = self.client.post("/api/study/quiz",json=self.request({"count":2}))
        self.assertEqual(result.status_code,200,result.text)
        return result.json()

    def report_response(self):
        return {"title":"酶活性实验", "sections":[{"key":key,"title":title,"content":f"固定内容：{title}","sources":[self.ref.copy()]} for key,title in [
            ("purpose","实验目的"),("principle","实验原理"),("equipment","仪器与试剂"),("steps","实验步骤"),("precautions","注意事项")]], "pre_questions":[self.question() for _ in range(3)]}

    def test_health_and_frontend(self):
        self.assertEqual(self.client.get("/api/health").json()["status"],"ok")
        response = self.client.get("/")
        self.assertEqual(response.status_code,200)
        self.assertIn("小家伙的小家伙",response.text)
        self.assertIn("frame-ancestors 'none'",response.headers["content-security-policy"])
        self.assertEqual(len(self.client.get("/api/courses").json()),4)

    def test_key_is_not_persisted_or_returned(self):
        settings = self.client.get("/api/settings").json()
        self.assertTrue(settings["key_present"])
        self.assertNotIn("api_key",settings)
        with db.connect() as con:
            saved = con.execute("SELECT data FROM settings").fetchone()[0]
        self.assertNotIn("test-only-key",saved)
        ai.API_KEY = ""
        self.assertFalse(self.client.get("/api/settings").json()["key_present"])
        self.assertEqual(self.client.post("/api/settings/test").status_code,428)

    def test_guide_documents_and_seekable_demo(self):
        html=self.client.get('/guide/使用说明书.html')
        self.assertEqual(html.status_code,200)
        self.assertIn('医学生零基础使用说明书',html.text)
        self.assertIn('script-src',html.headers['content-security-policy'])
        self.assertNotIn('onclick=',html.text)
        self.assertEqual(self.client.get('/guide/guide.js').status_code,200)
        pdf=self.client.get('/guide/使用说明书.pdf')
        self.assertEqual(pdf.status_code,200)
        self.assertTrue(pdf.content.startswith(b'%PDF'))
        for name in ['小家伙的小家伙-操作演示.mp4','小家伙的小家伙-小猫互动演示.mp4']:
            clip=self.client.get('/guide/media/'+name,headers={'Range':'bytes=0-31'})
            self.assertEqual(clip.status_code,206)
            self.assertEqual(len(clip.content),32)
            self.assertIn('video/mp4',clip.headers['content-type'])
            self.assertTrue(clip.headers['content-range'].startswith('bytes 0-31/'))
        player=self.client.get('/guide/小猫互动演示.html')
        self.assertEqual(player.status_code,200)
        self.assertIn('<video',player.text)
        self.assertNotIn('onclick=',player.text)
        self.assertEqual(self.client.get('/guide/kitten-demo.js').status_code,200)
        chapters=self.client.get('/guide/media/小猫互动演示章节.json').json()
        self.assertFalse(chapters['uses_model_api'])
        self.assertEqual(len(chapters['chapters']),10)
        starts=[chapter['start'] for chapter in chapters['chapters']]
        self.assertEqual(starts,sorted(starts))
        self.assertLess(starts[-1],chapters['duration'])

    def test_api_request_contract_and_test_connection(self):
        self.response = "OK"
        result = self.client.post("/api/settings/test")
        self.assertEqual(result.status_code,200)
        self.assertTrue(result.json()["generation_verified"])
        url,headers,body = self.captured[-1]
        self.assertEqual(str(url),"https://example.com/v1/chat/completions")
        self.assertEqual(headers["authorization"],"Bearer test-only-key")
        self.assertEqual(body["model"],"test-model")
        self.assertEqual(body["messages"], [{"role":"user", "content":"Reply with OK only."}])
        self.assertEqual(body["max_tokens"],128)

    def test_connection_distinguishes_truncation_from_generation(self):
        for content in [None, "", "O"]:
            with self.subTest(content=content):
                self.completion_body = {"choices":[{"finish_reason":"length", "message":{
                    "content":content, "reasoning_content":"private reasoning must not be returned"}}]}
                result = self.client.post("/api/settings/test")
                self.assertEqual(result.status_code,200,result.text)
                self.assertTrue(result.json()["ok"])
                self.assertFalse(result.json()["generation_verified"])
                self.assertIn("尚未验证完整内容生成",result.json()["message"])
                self.assertNotIn("private reasoning",result.text)
                self.assertNotIn("test-only-key",result.text)

    def test_truncated_generation_is_rejected_without_saving(self):
        self.completion_body = {"choices":[{"finish_reason":"length", "message":{"content":"{\"summary\":\"unfinished"}}]}
        result = self.client.post("/api/study/outline",json=self.request())
        self.assertEqual(result.status_code,502)
        self.assertIn("截断",result.json()["detail"])
        with db.connect() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM artifacts").fetchone()[0],0)
        self.assertEqual([r[2]["max_tokens"] for r in self.captured],[8192,16384])

    def test_outline_and_chat_recover_from_reasoning_truncation(self):
        self.client.put('/api/settings',json={'base_url':'https://api.deepseek.com','model':'user-selected-model'})
        outputs = [('/api/study/outline', {}, {
            'summary':'酶活性实验', 'items':[{'term':'温度','explanation':'酶活性受温度影响','importance':'核心',
            'recall_question':'如何控制变量？','sources':[self.ref.copy()]}], 'connections':[]}),
            ('/api/study/chat', {'message':'为什么保持其他条件一致？'}, {
            'answer':'便于比较温度的影响。','follow_up':'你会如何设置对照？','sources':[self.ref.copy()]})]
        for endpoint, extra, output in outputs:
            with self.subTest(endpoint=endpoint):
                self.captured.clear()
                self.completion_queue = [self.completion(None,'length'),self.completion(json.dumps(output,ensure_ascii=False))]
                result = self.client.post(endpoint,json=self.request(extra))
                self.assertEqual(result.status_code,200,result.text)
                self.assertNotIn('private reasoning',result.text)
                self.assertEqual([r[2]['max_tokens'] for r in self.captured],[8192,16384])
                for _,_,body in self.captured:
                    self.assertEqual(body['thinking'],{'type':'disabled'})
                    self.assertEqual(body['response_format'],{'type':'json_object'})
                    self.assertNotIn('temperature',body)
                self.assertEqual(self.captured[0][2]['messages'],self.captured[1][2]['messages'])
        with db.connect() as con:
            rows = con.execute('SELECT data FROM artifacts').fetchall()
        self.assertEqual(len(rows),2)
        self.assertNotIn('private reasoning',''.join(r[0] for r in rows))

    def test_output_settings_persist_and_old_settings_gain_defaults(self):
        with db.connect() as con:
            con.execute('UPDATE settings SET data=? WHERE id=1',(json.dumps({'base_url':'https://api.deepseek.com','model':'existing-model'}),))
        settings = self.client.get('/api/settings').json()
        self.assertEqual(settings['model'],'existing-model')
        self.assertEqual(settings['max_output_tokens'],8192)
        self.assertEqual(settings['deepseek_mode'],'fast')
        self.assertTrue(settings['key_present'])
        result = self.client.put('/api/settings',json={**settings,'max_output_tokens':16000,'deepseek_mode':'thinking'})
        self.assertEqual(result.status_code,200)
        self.assertEqual(self.client.get('/api/settings').json()['max_output_tokens'],16000)
        self.assertEqual(self.client.get('/api/settings').json()['deepseek_mode'],'thinking')
        for values in [{'max_output_tokens':511},{'max_output_tokens':65537},{'deepseek_mode':'invalid'}]:
            self.assertEqual(self.client.put('/api/settings',json={**settings,**values}).status_code,422)

    def test_deepseek_modes_and_extensions_are_limited_to_official_host(self):
        self.response = 'OK'
        for mode, expected in [('fast',{'type':'disabled'}),('thinking',{'type':'enabled'}),('default',None)]:
            self.client.put('/api/settings',json={'base_url':'https://api.deepseek.com/v1','model':'chosen-model','deepseek_mode':mode})
            self.assertEqual(self.client.post('/api/settings/test').status_code,200)
            self.assertEqual(self.captured[-1][2].get('thinking'),expected)
            self.assertNotIn('response_format',self.captured[-1][2])
        for host in ['example.com','api.deepseek.com.example.com']:
            self.client.put('/api/settings',json={'base_url':f'https://{host}/v1','model':'chosen-model','deepseek_mode':'thinking'})
            self.assertEqual(self.client.post('/api/settings/test').status_code,200)
            self.assertNotIn('thinking',self.captured[-1][2])
            self.assertNotIn('response_format',self.captured[-1][2])
            self.assertEqual(self.captured[-1][2]['temperature'],0.2)

    def test_retry_budget_is_bounded_and_test_connection_does_not_retry(self):
        self.completion_body = self.completion('', 'length')
        for budget, expected in [(40000,[40000,65536]),(65536,[65536])]:
            self.client.put('/api/settings',json={'base_url':'https://example.com/v1','model':'chosen','max_output_tokens':budget})
            self.captured.clear()
            result = self.client.post('/api/study/chat',json=self.request({'message':'温度的作用？'}))
            self.assertEqual(result.status_code,502)
            self.assertIn('未保存',result.json()['detail'])
            self.assertEqual([r[2]['max_tokens'] for r in self.captured],expected)
        self.captured.clear()
        self.assertEqual(self.client.post('/api/settings/test').status_code,200)
        self.assertEqual([r[2]['max_tokens'] for r in self.captured],[128])

    def test_no_generation_retry_on_api_error_or_invalid_json(self):
        for status in [400,401,429]:
            self.status = status
            self.captured.clear()
            self.assertEqual(self.client.post('/api/study/outline',json=self.request()).status_code,502)
            self.assertEqual(len(self.captured),1)
        self.status = 200
        self.response = 'not JSON'
        self.captured.clear()
        self.assertEqual(self.client.post('/api/study/outline',json=self.request()).status_code,502)
        self.assertEqual(len(self.captured),1)

    def test_connection_rejects_malformed_or_empty_responses(self):
        for body in [[], {}, {"choices":[]}, {"choices":[None]}, {"choices":[{"message":[]}]},
                     {"choices":[{"finish_reason":"length","message":{}}]},
                     {"choices":[{"finish_reason":"stop","message":{"content":""}}]},
                     {"choices":[{"finish_reason":"stop","message":{"content":None,"reasoning_content":"thinking only"}}]}]:
            with self.subTest(body=body):
                self.completion_body = body
                self.assertEqual(self.client.post("/api/settings/test").status_code,502)

    def test_invalid_api_addresses_and_clear_key(self):
        for base in ["http://example.com/v1", "https://user:password@example.com/v1", "https://example.com/v1/chat/completions", "https://example.com:invalid/v1", "https://[invalid/v1"]:
            response=self.client.put("/api/settings",json={"base_url":base,"model":"test"})
            self.assertEqual(response.status_code,400)
        self.assertEqual(self.client.put("/api/settings",json={"base_url":"http://127.0.0.1:11434/v1","model":"local","clear_key":True}).status_code,200)
        self.assertFalse(ai.API_KEY)

    def test_upstream_errors_are_redacted(self):
        self.response="should not be shown"
        for status in [401,403,404,429]:
            with self.subTest(status=status):
                self.status=status
                result=self.client.post("/api/settings/test")
                self.assertEqual(result.status_code,502)
                self.assertNotIn("should not be shown",result.text)
                self.assertNotIn("test-only-key",result.text)

    def test_reject_cross_origin_and_untrusted_host(self):
        response=self.client.post("/api/settings/test",headers={"Origin":"https://evil.example"})
        self.assertEqual(response.status_code,403)
        self.assertEqual(self.client.get("/api/health",headers={"Host":"evil.example"}).status_code,400)

    def test_pdf_and_scanned_pdf(self):
        document=fitz.open();page=document.new_page();page.insert_text((72,72),"Temperature controls enzyme activity.")
        result=self.upload("lecture.pdf",document.tobytes())
        self.assertEqual(result.status_code,200,result.text)
        details=self.client.get('/api/materials/'+result.json()['id']).json()
        self.assertEqual(details['chunks'][0]['unit'],'第 1 页')
        document=fitz.open();document.new_page()
        self.assertEqual(self.upload("scan.pdf",document.tobytes()).status_code,400)

    def test_pptx_text_and_notes(self):
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text="Cell injury"
        slide.placeholders[1].text="Necrosis and apoptosis"
        slide.notes_slide.notes_text_frame.text="Teaching notes"
        buffer=io.BytesIO();ppt.save(buffer)
        response=self.upload("lecture.pptx",buffer.getvalue(),course="pathology")
        self.assertEqual(response.status_code,200,response.text)
        chunks=self.client.get('/api/materials/'+response.json()['id']).json()['chunks']
        self.assertEqual(chunks[0]['unit'],'第 1 张幻灯片')
        self.assertIn('Teaching notes',chunks[0]['text'])

    def test_docx_and_gb_text(self):
        doc=Document();doc.add_paragraph('实验目的');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='试剂';table.cell(0,1).text='酶'
        buffer=io.BytesIO();doc.save(buffer)
        response=self.upload('experiment.docx',buffer.getvalue(),kind='experiment')
        self.assertEqual(response.status_code,200,response.text)
        chunks=self.client.get('/api/materials/'+response.json()['id']).json()['chunks']
        self.assertEqual(chunks[-1]['unit'],'表 1')
        self.assertIn('酶',chunks[-1]['text'])
        response=self.upload('old.txt','酶活性'.encode('gb18030'))
        self.assertEqual(response.status_code,200,response.text)

    def test_invalid_uploads_leave_no_orphans(self):
        before=list((db.DATA/'uploads').iterdir())
        for name,data in [('bad.docx',b'not a zip'),('empty.txt',b''),('old.ppt',b'no')]:
            self.assertEqual(self.upload(name,data).status_code,400)
        self.assertEqual(before,list((db.DATA/'uploads').iterdir()))

    def test_course_isolation_and_template_not_evidence(self):
        response=self.client.post('/api/study/outline',json={'course':'pathology','material_ids':[self.material]})
        self.assertEqual(response.status_code,400)
        template=self.upload('template.txt','实验目的\n实验结果'.encode(),kind='template').json()['id']
        response=self.client.post('/api/study/outline',json=self.request({'material_ids':[template]}))
        self.assertEqual(response.status_code,400)
        self.assertFalse(self.captured)

    def test_outline_citations_and_idempotent_recall(self):
        self.response={'summary':'酶活性实验','items':[{'term':'温度','explanation':'酶活性受温度影响','importance':'核心','recall_question':'实验为何保持其他条件一致？','sources':[self.ref.copy()]}], 'connections':[{'chain':'温度 → 酶活性','explanation':'本实验比较温度的作用','sources':[self.ref.copy()]}]}
        result=self.client.post('/api/study/outline',json=self.request())
        self.assertEqual(result.status_code,200,result.text)
        data=result.json();self.assertEqual(data['data']['items'][0]['sources'][0]['title'],'实验.txt')
        self.assertEqual(self.client.post('/api/artifacts/'+data['id']+'/recall').json()['added'],1)
        self.assertEqual(self.client.post('/api/artifacts/'+data['id']+'/recall').json()['added'],0)
        self.assertEqual(len(self.client.get('/api/reviews?course=biochemistry').json()),1)
        self.assertIn('上传材料和学生陈述',self.captured[-1][2]['messages'][0]['content'])

    def test_false_citation_is_rejected_before_saving(self):
        self.response={'answer':'错误引用','follow_up':'','sources':[{'material_id':self.material,'unit':'第 99 页'}]}
        response=self.client.post('/api/study/chat',json=self.request({'message':'解释酶活性'}))
        self.assertEqual(response.status_code,502)
        self.assertEqual(len(self.captured),2)
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_short_citations_resolve_to_selected_pdf_pages(self):
        document = fitz.open()
        for text in ['Signal transduction: receptors.','Signal transduction: second messengers.']:
            document.new_page().insert_text((72,72),text)
        uploaded = self.upload('信号转导.pdf',document.tobytes()).json()['id']
        self.response = self.outline_response([{'source_id':' S1 '}])
        self.response['connections'] = [{'chain':'受体 → 信号','explanation':'材料中的信号通路','sources':[{'source_id':'S2'}]}]
        result = self.client.post('/api/study/outline',json=self.request({'material_ids':[uploaded]}))
        self.assertEqual(result.status_code,200,result.text)
        data = result.json()['data']
        for source, page in [(data['items'][0]['sources'][0],1),(data['connections'][0]['sources'][0],2)]:
            self.assertEqual(source,{'material_id':uploaded,'unit':f'第 {page} 页','title':'信号转导.pdf'})
        self.assertNotIn('source_id',json.dumps(data))
        prompt = self.captured[0][2]['messages'][1]['content']
        self.assertNotIn(self.material,prompt)
        self.assertIn('[来源 S1]\n[材料 '+uploaded,prompt)
        self.assertIn('"enum": ["S1", "S2"]',prompt)
        self.assertIn('"required": ["source_id"]',prompt)

    def test_unknown_short_citation_is_regenerated_with_valid_sources(self):
        wrong = self.outline_response([{'source_id':'S999'}])
        wrong['summary'] = 'invalid-answer-not-to-be-echoed'
        correct = self.outline_response([{'source_id':'S1'}])
        self.completion_queue = [self.completion(json.dumps(wrong)),self.completion(json.dumps(correct))]
        result = self.client.post('/api/study/outline',json=self.request())
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['data']['items'][0]['sources'][0]['material_id'],self.material)
        self.assertEqual(len(self.captured),2)
        repair = self.captured[-1][2]['messages'][1]['content']
        self.assertIn('来源校正重试',repair)
        self.assertNotIn('invalid-answer-not-to-be-echoed',repair)
        with db.connect() as con:
            self.assertEqual(con.execute('SELECT COUNT(*) FROM artifacts').fetchone()[0],1)

    def test_unknown_or_mixed_short_citations_are_never_saved(self):
        for source in [{'source_id':'S999'}, {'source_id':None},
                       {'source_id':'S1','material_id':'unselected','unit':'第 99 页'}]:
            with self.subTest(source=source):
                self.captured.clear()
                self.response = self.outline_response([source])
                result = self.client.post('/api/study/outline',json=self.request())
                self.assertEqual(result.status_code,502)
                self.assertIn('来源',result.json()['detail'])
                self.assertEqual(len(self.captured),2)
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_sources_in_unselected_material_are_rejected(self):
        other = self.upload('其他材料.txt','其他章节'.encode()).json()['id']
        self.response = self.outline_response([{'material_id':other,'unit':'段落 1'}])
        result = self.client.post('/api/study/outline',json=self.request())
        self.assertEqual(result.status_code,502)
        self.assertEqual(len(self.captured),2)
        self.assertNotIn(other,self.captured[0][2]['messages'][1]['content'])
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_unretrieved_pdf_page_is_rejected_even_in_selected_file(self):
        document = fitz.open()
        for index in range(40):
            document.new_page().insert_text((72,72),f'Lecture page {index+1}: signal transduction.')
        uploaded = self.upload('long.pdf',document.tobytes()).json()['id']
        self.response = self.outline_response([{'material_id':uploaded,'unit':'第 40 页'}])
        result = self.client.post('/api/study/outline',json=self.request({'material_ids':[uploaded]}))
        self.assertEqual(result.status_code,502)
        self.assertEqual(len(self.captured),2)
        self.assertNotIn('第 40 页',self.captured[0][2]['messages'][1]['content'])
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_repeated_chunks_of_same_page_share_citation_code(self):
        chunks = [{'unit':'第 1 页','text':'受体与配体'}, {'unit':'第 1 页','text':'受体的下游效应'},
                  {'unit':'第 2 页','text':'第二信使'}]
        with db.connect() as con:
            con.execute('UPDATE materials SET chunks=? WHERE id=?',(json.dumps(chunks),self.material))
        self.response = self.outline_response([{'source_id':'S2'}])
        result = self.client.post('/api/study/outline',json=self.request())
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()['data']['items'][0]['sources'][0]['unit'],'第 2 页')
        prompt = self.captured[0][2]['messages'][1]['content']
        self.assertEqual(prompt.count('[来源 S1]'),2)
        self.assertIn('"enum": ["S1", "S2"]',prompt)

    def test_citation_repair_and_length_share_single_retry(self):
        bad = json.dumps(self.outline_response([{'source_id':'S999'}]))
        for completions in [[self.completion(None,'length'),self.completion(bad)],
                            [self.completion(bad),self.completion(None,'length')]]:
            with self.subTest(completions=completions):
                self.captured.clear()
                self.completion_queue = completions
                result = self.client.post('/api/study/outline',json=self.request())
                self.assertEqual(result.status_code,502)
                self.assertEqual(len(self.captured),2)
                self.assertEqual(self.completion_queue,[])
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_malformed_model_output_is_rejected(self):
        self.response='not JSON'
        response=self.client.post('/api/study/outline',json=self.request())
        self.assertEqual(response.status_code,502)
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_teach_and_chat_history(self):
        self.response={'understood':['提到了温度'],'corrections':[{'student_claim':'所有条件同时变化','feedback':'其他条件应一致','sources':[self.ref.copy()]}],'missing':[],'next_question':'为什么需要对照组？'}
        response=self.client.post('/api/study/feedback',json=self.request({'topic':'实验设计','explanation':'所有条件都同时变化。'}))
        self.assertEqual(response.status_code,200,response.text)
        self.response={'answer':'保持其他条件一致','follow_up':'你会如何设置对照？','sources':[self.ref.copy()]}
        response=self.client.post('/api/study/chat',json=self.request({'message':'为什么？','history':[{'role':'user','content':'对照组是什么？'}]}))
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(response.json()['data']['question'],'为什么？')
        self.assertIn('对照组是什么',self.captured[-1][2]['messages'][1]['content'])

    def test_quiz_answers_hidden_grade_and_review(self):
        quiz=self.create_quiz()
        self.assertNotIn('answer',quiz['data']['questions'][0])
        self.assertNotIn('explanation',self.client.get('/api/artifacts/'+quiz['id']).json()['data']['questions'][0])
        invalid=self.client.post('/api/artifacts/'+quiz['id']+'/grade',json={'answers':[0,None]})
        self.assertEqual(invalid.status_code,400)
        grade=self.client.post('/api/artifacts/'+quiz['id']+'/grade',json={'answers':[0,2]}).json()
        self.assertEqual(grade['correct'],1);self.assertEqual(grade['total'],2)
        cards=self.client.get('/api/reviews?course=biochemistry').json();self.assertEqual(len(cards),1)
        self.assertNotIn('answer',cards[0])
        card=cards[0]['id']
        self.assertEqual(self.client.post('/api/reviews/'+card+'/answer',json={'answer':''}).status_code,400)
        revealed=self.client.post('/api/reviews/'+card+'/answer',json={'answer':'其他条件'}).json()
        self.assertEqual(revealed['answer'],'其他条件')
        result=self.client.post('/api/reviews/'+card+'/answer',json={'answer':'其他条件','remembered':True}).json()
        due=datetime.fromisoformat(result['due'])
        self.assertAlmostEqual((due-datetime.now(timezone.utc)).total_seconds()/86400,1,places=2)
        self.assertEqual(self.client.get('/api/reviews?course=biochemistry').json(),[])

    def test_quiz_wrong_number_rejected(self):
        self.response={'questions':[self.question()]}
        self.assertEqual(self.client.post('/api/study/quiz',json=self.request({'count':2})).status_code,502)

    def test_comparison_practice(self):
        self.response={'left_name':'A','right_name':'B','rows':[{'dimension':'作用','left':'A作用','right':'B作用','sources':[self.ref.copy()]},{'dimension':'条件','left':'A条件','right':'B条件','sources':[self.ref.copy()]}],'memory_tip':'按条件区分','practice':[self.question()]}
        response=self.client.post('/api/study/compare',json=self.request({'left':'A','right':'B'}))
        self.assertEqual(response.status_code,200,response.text)
        data=response.json();self.assertNotIn('answer',data['data']['practice'][0])
        response=self.client.post('/api/artifacts/'+data['id']+'/grade',json={'answers':[0]})
        self.assertEqual(response.json()['correct'],1)

    def test_report_template_export_and_linked_quiz(self):
        doc=Document();p=doc.add_paragraph();p.add_run('{{实验');p.add_run('名称}}')
        table=doc.add_table(rows=2,cols=1);table.cell(0,0).text='{{实验目的}}';table.cell(1,0).text='{{实验数据}}'
        buffer=io.BytesIO();doc.save(buffer)
        template=self.upload('school.docx',buffer.getvalue(),kind='template').json()['id']
        self.response=self.report_response()
        response=self.client.post('/api/reports',json=self.request({'title':'酶活性实验','template_id':template}))
        self.assertEqual(response.status_code,200,response.text)
        data=response.json()['data'];export=self.client.get('/api/exports/'+data['export_id'])
        self.assertEqual(export.status_code,200)
        rendered=Document(io.BytesIO(export.content))
        self.assertEqual(rendered.paragraphs[0].text,'酶活性实验')
        self.assertEqual(rendered.tables[0].cell(0,0).text,'固定内容：实验目的')
        self.assertIn('真实实验数据',rendered.tables[0].cell(1,0).text)
        all_text='\n'.join(p.text for p in rendered.paragraphs)
        self.assertIn('根据真实结果填写结论',all_text)
        self.assertIn('实验.txt',all_text)
        quiz=self.client.get('/api/artifacts/'+data['quiz_id']).json()
        self.assertEqual(len(quiz['data']['questions']),3)
        self.assertNotIn('answer',quiz['data']['questions'][0])

    def test_report_rejects_duplicate_fixed_sections(self):
        self.response=self.report_response();self.response['sections'][0]['key']='steps'
        response=self.client.post('/api/reports',json=self.request({'title':'实验'}))
        self.assertEqual(response.status_code,502)
        self.assertEqual(list((db.DATA/'exports').iterdir()),[])

    def test_headings_template_removes_example_conclusions(self):
        doc=Document();doc.add_paragraph('姓名：____');doc.add_heading('一、实验目的',1);doc.add_paragraph('旧目的');doc.add_heading('实验结论',1);doc.add_paragraph('样例结论：全部成功')
        buffer=io.BytesIO();doc.save(buffer)
        template=self.upload('school.docx',buffer.getvalue(),kind='template').json()['id']
        self.response=self.report_response()
        data=self.client.post('/api/reports',json=self.request({'title':'实验','template_id':template})).json()['data']
        rendered=Document(io.BytesIO(self.client.get('/api/exports/'+data['export_id']).content))
        text='\n'.join(p.text for p in rendered.paragraphs)
        self.assertNotIn('全部成功',text);self.assertNotIn('旧目的',text)
        self.assertIn('姓名：____',text);self.assertIn('固定内容：实验目的',text)

    def test_figure_labels_hidden_grade_aliases_and_review_cue(self):
        image=Image.new('RGB',(200,100),'white');buffer=io.BytesIO();image.save(buffer,'PNG')
        response=self.client.post('/api/figures',data={'course':'cell-biology'},files={'file':('cell.png',buffer.getvalue())})
        id=response.json()['id']
        self.assertEqual(self.client.get('/api/figures/'+id+'/practice').status_code,400)
        labels=[{'id':'p1','name':'细胞核 / nucleus','x':0.5,'y':0.5,'explanation':'根据教材标注核的位置。'},{'id':'p2','name':'细胞膜','x':0.8,'y':0.2,'explanation':'根据教材标注细胞边界。'}]
        response=self.client.put('/api/figures/'+id+'/labels',json={'labels':labels,'confirmed':True})
        self.assertEqual(response.status_code,200)
        practice=self.client.get('/api/figures/'+id+'/practice').json()
        self.assertNotIn('name',practice['points'][0])
        self.assertNotIn('explanation',practice['points'][0])
        grade=self.client.post('/api/figures/'+id+'/grade',json={'answers':{'p1':'Nucleus','p2':'错误'}}).json()
        self.assertEqual(grade['correct'],1)
        cards=self.client.get('/api/reviews?course=cell-biology').json()
        self.assertEqual(cards[0]['figure']['x'],0.8)
        self.assertNotIn('answer',cards[0])

    def test_deleting_material_preserves_artifacts(self):
        quiz=self.create_quiz()
        self.assertEqual(self.client.delete('/api/materials/'+self.material).status_code,200)
        self.assertEqual(self.client.get('/api/materials/'+self.material).status_code,404)
        self.assertEqual(self.client.get('/api/artifacts/'+quiz['id']).status_code,200)
        self.assertEqual(list((db.DATA/'uploads').iterdir()),[])


    def chapter_document(self):
        doc = Document()
        for index, name in enumerate(['第一章 细胞膜','第二章 信号转导','第三章 细胞周期'],1):
            doc.add_heading(name,1)
            doc.add_paragraph(f'CHAPTER_{index}_ONLY 这是第 {index} 章的独立内容。')
            doc.add_table(rows=1,cols=1).cell(0,0).text=f'TABLE_{index}_ONLY'
        buf=io.BytesIO();doc.save(buf)
        return self.upload('细胞生物学.docx',buf.getvalue()).json()['id']

    def agent_plan(self, tasks=None, message='我会按指定章节完成任务。', status='execute'):
        return self.completion(json.dumps({'status':status,'message':message,'tasks':tasks or []},ensure_ascii=False))

    def test_docx_chapter_scope_keeps_tables_in_their_original_chapter(self):
        material=self.chapter_document()
        catalog=self.client.get(f'/api/materials/{material}/chapters').json()
        self.assertEqual([c['title'] for c in catalog['chapters']],['第一章 细胞膜','第二章 信号转导','第三章 细胞周期'])
        self.response=self.outline_response([{'source_id':'S1'}])
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material], 'scopes':[{'material_id':material,'chapter_id':'C2'}]}))
        self.assertEqual(result.status_code,200,result.text)
        prompt=self.captured[-1][2]['messages'][1]['content']
        self.assertIn('CHAPTER_2_ONLY',prompt);self.assertIn('TABLE_2_ONLY',prompt)
        for token in ['CHAPTER_1_ONLY','CHAPTER_3_ONLY','TABLE_1_ONLY','TABLE_3_ONLY']:
            self.assertNotIn(token,prompt)
        self.assertEqual(result.json()['data']['evidence']['ranges'][0]['label'],'第二章 信号转导')

    def test_chapters_within_one_text_unit_are_sliced_before_generation(self):
        material=self.upload('chapters.txt','第一章 基础\nFIRST_ONLY\n第二章 信号\nSECOND_ONLY\n第三章 周期\nTHIRD_ONLY'.encode()).json()['id']
        self.response=self.outline_response([{'source_id':'S1'}])
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':'C2'}]}))
        self.assertEqual(result.status_code,200,result.text)
        prompt=self.captured[-1][2]['messages'][1]['content']
        self.assertIn('SECOND_ONLY',prompt)
        self.assertNotIn('FIRST_ONLY',prompt);self.assertNotIn('THIRD_ONLY',prompt)

    def test_overlapping_text_chunks_do_not_leak_previous_chapter(self):
        text='第一章 基础\n'+('甲'*1300)+'FIRST_EDGE_ONLY\n第二章 信号\n'+('乙'*1500)+'SECOND_ONLY\n第三章 周期\nTHIRD_ONLY'
        material=self.upload('overlap.txt',text.encode()).json()['id']
        self.response=self.outline_response([{'source_id':'S1'}])
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':'C2'}]}))
        self.assertEqual(result.status_code,200,result.text)
        prompt=self.captured[-1][2]['messages'][1]['content']
        self.assertIn('SECOND_ONLY',prompt)
        self.assertNotIn('FIRST_EDGE_ONLY',prompt);self.assertNotIn('THIRD_ONLY',prompt)

    def test_pdf_toc_and_manual_page_range_are_exact(self):
        doc=fitz.open()
        for index in range(4):doc.new_page().insert_text((72,72),f'PAGE_{index+1}_ONLY')
        doc.set_toc([[1,'第一章',1],[1,'第二章',3]])
        material=self.upload('目录.pdf',doc.tobytes()).json()['id']
        catalog=self.client.get(f'/api/materials/{material}/chapters').json()
        self.assertEqual(catalog['chapters'][1]['start_unit'],'第 3 页')
        self.response=self.outline_response([{'source_id':'S1'}])
        for spec, included, excluded in [({'chapter_id':'C2'},['PAGE_3_ONLY','PAGE_4_ONLY'],['PAGE_1_ONLY','PAGE_2_ONLY']),
                                        ({'start_unit':'第2页','end_unit':'第 3 页'},['PAGE_2_ONLY','PAGE_3_ONLY'],['PAGE_1_ONLY','PAGE_4_ONLY'])]:
            result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,**spec}]}))
            self.assertEqual(result.status_code,200,result.text)
            prompt=self.captured[-1][2]['messages'][1]['content']
            for text in included:self.assertIn(text,prompt)
            for text in excluded:self.assertNotIn(text,prompt)

    def test_invalid_scope_never_falls_back_to_all_materials(self):
        material=self.chapter_document()
        for scopes in [[{'material_id':material,'chapter_id':'C99'}], [{'material_id':self.material,'chapter_id':'C1'}],
                       [{'material_id':material,'start_unit':'第 6 段','end_unit':'第 1 段'}],
                       [{'material_id':material,'chapter_id':'C2','start_unit':'第 1 段'}]]:
            self.captured.clear()
            result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':scopes}))
            self.assertEqual(result.status_code,400,result.text)
            self.assertEqual(self.captured,[])

    def test_pdf_export_keeps_chinese_scope_sources_and_long_text(self):
        material=self.chapter_document()
        self.response=self.outline_response([{'source_id':'S1'}])
        self.response['items'][0]['explanation']='信号转导机制。'*450+'末尾哨兵'
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':'C2'}]})).json()
        before=len(self.captured)
        exported=self.client.get('/api/artifacts/'+result['id']+'/pdf')
        self.assertEqual(exported.status_code,200,exported.text[:100] if not exported.content.startswith(b'%PDF') else '')
        with fitz.open(stream=exported.content,filetype='pdf') as pdf:
            self.assertGreater(len(pdf),1)
            text=''.join(p.get_text() for p in pdf)
        for value in ['末尾哨兵','第二章 信号转导','细胞生物学.docx','主动回忆']:
            self.assertIn(value,text)
        self.assertNotIn('test-only-key',text)
        self.assertEqual(self.client.get('/api/artifacts/'+result['id']+'/pdf').content,exported.content)
        self.assertEqual(len(self.captured),before)

    def test_quiz_pdf_does_not_reveal_hidden_explanations(self):
        quiz=self.create_quiz()
        response=self.client.get('/api/artifacts/'+quiz['id']+'/pdf')
        self.assertEqual(response.status_code,200)
        with fitz.open(stream=response.content,filetype='pdf') as pdf:
            text=''.join(p.get_text() for p in pdf)
        self.assertIn('设计温度实验',text)
        self.assertNotIn('保持其他条件一致，才能比较温度的作用。',text)
        self.assertIn('练习版',text)

    def test_agent_clarifies_without_materials_or_reliable_chapters(self):
        result=self.client.post('/api/agent',json={'course':'biochemistry','material_ids':[],'message':'帮我学习'}).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        result=self.client.post('/api/agent',json=self.request({'message':'只整理第二章知识清单'})).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        self.assertIn('起止',result['messages'][-1]['content'])
        self.assertEqual(self.captured,[])

    def test_agent_clarifies_then_executes_two_tasks_in_one_chapter(self):
        material=self.chapter_document()
        self.completion_queue=[self.agent_plan(message='你想做知识清单还是小测？',status='clarify')]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'message':'帮我复习'})).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        tasks=[{'action':'outline','material_key':'M1','chapter_key':'C2','topic':'信号转导','instructions':'只列 1 个点'},
               {'action':'quiz','material_key':'M1','chapter_key':'C2','topic':'信号转导','count':2}]
        q=self.question();q['sources']=[{'source_id':'S1'}]
        self.completion_queue=[self.agent_plan(tasks),self.completion(json.dumps(self.outline_response([{'source_id':'S1'}]))),self.completion(json.dumps({'questions':[q,q]}))]
        response=self.client.post('/api/agent',json=self.request({'material_ids':[material],'thread_id':result['id'],'message':'只用第二章，先做清单，再出2道题'}))
        self.assertEqual(response.status_code,200,response.text)
        result=response.json();turn=result['messages'][-1]
        self.assertEqual(turn['status'],'completed')
        self.assertEqual(len(turn['results']),2)
        self.assertEqual(len(result['messages']),4)
        for request in self.captured[-2:]:
            prompt=request[2]['messages'][1]['content']
            self.assertIn('CHAPTER_2_ONLY',prompt);self.assertNotIn('CHAPTER_1_ONLY',prompt)
        self.assertEqual(self.client.get(f'/api/agent/threads/{result["id"]}?course=biochemistry').json()['messages'],result['messages'])
        for item in turn['results']:
            self.assertEqual(self.client.get(f'/api/artifacts/{item["artifact_id"]}/pdf').status_code,200)

    def test_agent_followup_preserves_scope_and_pdf_download_needs_no_model(self):
        material=self.chapter_document()
        self.completion_queue=[self.agent_plan([{'action':'outline','material_key':'M1','chapter_key':'C2'}]),self.completion(json.dumps(self.outline_response([{'source_id':'S1'}])))]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'message':'只用第二章做清单'})).json()
        self.completion_queue=[self.agent_plan([{'action':'chat','topic':'请解释刚才章节的机制'}]),self.completion(json.dumps({'answer':'解释信号转导。','follow_up':'如何调节？','sources':[{'source_id':'S1'}]}))]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'thread_id':result['id'],'message':'再解释一下机制'})).json()
        self.assertEqual(result['messages'][-1]['status'],'completed')
        self.assertNotIn('CHAPTER_1_ONLY',self.captured[-1][2]['messages'][1]['content'])
        calls=len(self.captured)
        exported=self.client.post('/api/agent',json=self.request({'material_ids':[material],'thread_id':result['id'],'message':'把刚才的结果下载 PDF'})).json()
        self.assertEqual(exported['messages'][-1]['results'][0]['artifact_id'],result['messages'][-1]['results'][0]['artifact_id'])
        self.assertEqual(len(self.captured),calls)

    def test_agent_invalid_tool_and_unknown_chapter_do_not_execute(self):
        material=self.chapter_document()
        for task, expected in [({'action':'delete_files'},502),({'action':'outline','material_key':'M1','chapter_key':'C99'},200)]:
            self.captured.clear()
            self.completion_queue=[self.agent_plan([task])]
            response=self.client.post('/api/agent',json=self.request({'material_ids':[material],'message':'请帮我处理'}))
            self.assertEqual(response.status_code,expected,response.text)
            self.assertEqual(len(self.captured),1)
        self.assertEqual(self.client.get('/api/artifacts?course=biochemistry').json(),[])

    def test_agent_threads_cannot_be_used_in_another_course(self):
        result=self.client.post('/api/agent',json={'course':'biochemistry','material_ids':[],'message':'帮我学习'}).json()
        self.assertEqual(self.client.get(f'/api/agent/threads/{result["id"]}?course=pathology').status_code,404)
        self.assertEqual(self.client.post('/api/agent',json={'course':'pathology','thread_id':result['id'],'message':'继续'}).status_code,404)
        self.assertEqual(self.captured,[])

    def test_parent_chapter_contains_its_subsections(self):
        material=self.upload('sections.txt','第一章 概述\nFIRST_ONLY\n第二章 信号\nINTRO_SECOND\n第一节 受体\nRECEPTOR_ONLY\n第二节 信使\nMESSENGER_ONLY\n第三章 周期\nTHIRD_ONLY'.encode()).json()['id']
        catalog=self.client.get(f'/api/materials/{material}/chapters').json()
        chapter=next(c for c in catalog['chapters'] if c['title']=='第二章 信号')
        self.response=self.outline_response([{'source_id':'S1'}])
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':chapter['id']}]}))
        self.assertEqual(result.status_code,200,result.text)
        prompt=self.captured[-1][2]['messages'][1]['content']
        for value in ['INTRO_SECOND','RECEPTOR_ONLY','MESSENGER_ONLY']:self.assertIn(value,prompt)
        for value in ['FIRST_ONLY','THIRD_ONLY']:self.assertNotIn(value,prompt)

    def test_pptx_chapters_use_slide_titles_not_chapter_mentions_in_notes(self):
        ppt=Presentation()
        for title,body in [('第一章 细胞膜','FIRST_ONLY'),('膜转运','FIRST_MORE'),('第二章 信号','SECOND_ONLY'),('信号放大','SECOND_MORE')]:
            slide=ppt.slides.add_slide(ppt.slide_layouts[1]);slide.shapes.title.text=title;slide.placeholders[1].text=body
            slide.notes_slide.notes_text_frame.text='第三章 这是备注中的参考章节，不是标题'
        buffer=io.BytesIO();ppt.save(buffer)
        material=self.upload('chapters.pptx',buffer.getvalue()).json()['id']
        catalog=self.client.get(f'/api/materials/{material}/chapters').json()
        self.assertEqual([c['title'] for c in catalog['chapters']],['第一章 细胞膜','第二章 信号'])
        self.response=self.outline_response([{'source_id':'S1'}])
        result=self.client.post('/api/study/outline',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':'C2'}]}))
        self.assertEqual(result.status_code,200,result.text)
        prompt=self.captured[-1][2]['messages'][1]['content']
        self.assertIn('SECOND_MORE',prompt);self.assertNotIn('FIRST_ONLY',prompt)

    def test_agent_can_clear_an_explicit_range_only_when_requested(self):
        material=self.chapter_document()
        spec={'material_id':material,'chapter_id':'C2'}
        self.completion_queue=[self.agent_plan([{'action':'outline','scope_mode':'all'}]),self.completion(json.dumps(self.outline_response([{'source_id':'S1'}])))]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'scopes':[spec],'message':'取消限制，改成整份材料的知识清单'})).json()
        self.assertEqual(result['messages'][-1]['status'],'completed')
        self.assertEqual(result['last_scopes'],[])
        self.assertIn('CHAPTER_1_ONLY',self.captured[-1][2]['messages'][1]['content'])
        self.captured.clear()
        self.completion_queue=[self.agent_plan([{'action':'outline','scope_mode':'all'}])]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'scopes':[spec],'message':'再整理一下'})).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        self.assertEqual(len(self.captured),1)

    def test_agent_ui_can_clear_saved_scope(self):
        material=self.chapter_document()
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'scopes':[{'material_id':material,'chapter_id':'C2'}],'message':'帮我复习'})).json()
        self.completion_queue=[self.agent_plan([{'action':'outline'}]),self.completion(json.dumps(self.outline_response([{'source_id':'S1'}])))]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'thread_id':result['id'],'reset_scope':True,'message':'整理知识清单'})).json()
        self.assertEqual(result['last_scopes'],[])
        self.assertEqual(result['messages'][-1]['status'],'completed')
        self.assertIn('CHAPTER_1_ONLY',self.captured[-1][2]['messages'][1]['content'])

    def test_agent_cannot_silently_expand_scope_by_switching_files(self):
        material=self.chapter_document()
        other=self.upload('另一份.txt',b'OTHER_FILE_ONLY').json()['id']
        other_key=f'M{sorted([material,other]).index(other)+1}'
        self.completion_queue=[self.agent_plan([{'action':'outline','material_key':other_key}])]
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material,other],'scopes':[{'material_id':material,'chapter_id':'C2'}],'message':'再做一个清单'})).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        self.assertEqual(len(self.captured),1)

    def test_vague_agent_request_preserves_explicit_scope_without_model_call(self):
        material=self.chapter_document();spec={'material_id':material,'chapter_id':'C2'}
        result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'scopes':[spec],'message':'帮我复习'})).json()
        self.assertEqual(result['messages'][-1]['status'],'clarify')
        self.assertEqual(result['last_scopes'][0]['chapter_id'],'C2')
        self.assertEqual(self.captured,[])

    def test_agent_keeps_successful_results_when_another_tool_fails(self):
        from fastapi import HTTPException
        from unittest.mock import AsyncMock
        material=self.chapter_document()
        tasks=[{'action':'outline','material_key':'M1','chapter_key':'C2'}, {'action':'quiz','material_key':'M1','chapter_key':'C2','count':2}]
        self.completion_queue=[self.agent_plan(tasks),self.completion(json.dumps(self.outline_response([{'source_id':'S1'}])))]
        with patch('app.main.quiz',new=AsyncMock(side_effect=HTTPException(502,'模拟服务暂时不可用'))):
            result=self.client.post('/api/agent',json=self.request({'material_ids':[material],'message':'只用第二章做清单和小测'})).json()
        self.assertEqual(result['messages'][-1]['status'],'partial')
        results=result['messages'][-1]['results']
        self.assertEqual([r['status'] for r in results],['completed','error'])
        self.assertEqual(self.client.get('/api/artifacts/'+results[0]['artifact_id']).status_code,200)

if __name__=='__main__':
    unittest.main()
